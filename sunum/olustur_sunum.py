#!/usr/bin/env python3
"""Endüstriyel tesisler için Mavi Su Verimliliği Belgesi sunumu üreticisi."""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent
IMG = ROOT / "gorseller"
OUT = ROOT / "Endustriyel_Tesisler_Mavi_Su_Belgesi.pptx"
CACHE = ROOT / "_cache"
CACHE.mkdir(exist_ok=True)

# Widescreen 16:9
W, H = Inches(13.333), Inches(7.5)

NAVY = RGBColor(10, 36, 72)
NAVY2 = RGBColor(15, 52, 96)
TEAL = RGBColor(0, 148, 154)
TEAL_D = RGBColor(0, 112, 118)
GOLD = RGBColor(214, 148, 42)
AMBER = RGBColor(196, 112, 28)
CREAM = RGBColor(247, 242, 232)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(232, 244, 247)
MINT = RGBColor(230, 246, 238)
SAND = RGBColor(255, 244, 224)
LAV = RGBColor(236, 232, 252)
INK = RGBColor(28, 39, 54)
MUTED = RGBColor(90, 104, 122)
GREEN = RGBColor(26, 138, 99)
CORAL = RGBColor(214, 78, 72)
SKY = RGBColor(36, 120, 196)
PURPLE = RGBColor(92, 78, 168)

FONT = "Calibri"


def emu(inches: float) -> int:
    return int(Inches(inches))


def set_run(run, size, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", FONT)


def add_text(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, italic)
    return p


def add_para(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, space_before=0, space_after=4):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def fill_line(shape, color, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()


def rect(slide, l, t, w, h, color, rounded=True, radius=0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h),
    )
    fill(shp, color)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def pill(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    fill(shp, color)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    return shp


def oval(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(w), Inches(h))
    fill(shp, color)
    return shp


def tb(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    box.text_frame.paragraphs[0].alignment = align
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    add_text(tf, text, size, bold, color, align)
    return box


def tb_multi(slide, l, t, w, h, lines, anchor=MSO_ANCHOR.TOP):
    """lines: list of (text, size, bold, color, space_before)."""
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    first = True
    for item in lines:
        text, size = item[0], item[1]
        bold = item[2] if len(item) > 2 else False
        color = item[3] if len(item) > 3 else INK
        spb = item[4] if len(item) > 4 else 0
        if first:
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(spb)
            run = p.add_run()
            run.text = text
            set_run(run, size, bold, color)
            first = False
        else:
            add_para(tf, text, size, bold, color, space_before=spb, space_after=2)
    return box


def picture(slide, path, l, t, w, h):
    slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h))


def crop_cover(src: Path, dest: Path, size=(1920, 1080), dark_left=True):
    im = Image.open(src).convert("RGB")
    im = ImageOps.fit(im, size, Image.Resampling.LANCZOS)
    if dark_left:
        overlay = Image.new("RGB", size, (8, 28, 56))
        mask = Image.new("L", size, 0)
        d = ImageDraw.Draw(mask)
        for x in range(size[0]):
            if x < int(size[0] * 0.52):
                a = 185
            else:
                a = max(0, int(185 - (x - size[0] * 0.52) / (size[0] * 0.28) * 185))
            d.line([(x, 0), (x, size[1])], fill=a)
        im = Image.composite(overlay, im, mask)
        im = ImageEnhance.Color(im).enhance(1.08)
    im.save(dest, quality=92)
    return dest


def crop_fit(src: Path, dest: Path, size, boost=1.08):
    im = Image.open(src).convert("RGB")
    im = ImageOps.fit(im, size, Image.Resampling.LANCZOS)
    im = ImageEnhance.Color(im).enhance(boost)
    im = ImageEnhance.Contrast(im).enhance(1.04)
    im.save(dest, quality=90)
    return dest


def footer(slide, n, total=13, dark=False):
    color = RGBColor(180, 198, 214) if dark else MUTED
    tb(slide, 0.4, 7.18, 9.5, 0.25,
       "Endüstriyel Tesisler için Mavi Su Verimliliği Belgesi  •  Tarım ve Orman Bakanlığı mevzuatı esas alınmıştır",
       10, False, color)
    tb(slide, 12.2, 7.18, 0.8, 0.25, f"{n}/{total}", 10, True, color, PP_ALIGN.RIGHT)


def tr_upper(text: str) -> str:
    return text.translate(str.maketrans({"i": "İ", "ı": "I"})).upper()


def header_bar(slide, kicker, title, subtitle=None):
    rect(slide, 0, 0, 13.333, 0.12, GOLD, rounded=False)
    rect(slide, 0, 0.12, 13.333, 1.18, NAVY, rounded=False)
    tb(slide, 0.45, 0.22, 12.4, 0.28, tr_upper(kicker), 11, True, GOLD)
    tb(slide, 0.45, 0.48, 12.4, 0.42, title, 24, True, WHITE)
    if subtitle:
        tb(slide, 0.45, 0.92, 12.4, 0.28, subtitle, 12, False, RGBColor(186, 214, 230))


def card(slide, l, t, w, h, color=WHITE):
    shadow = rect(slide, l + 0.04, t + 0.05, w, h, RGBColor(220, 228, 234), radius=0.08)
    return rect(slide, l, t, w, h, color, radius=0.08)


def num_badge(slide, l, t, n, color=TEAL):
    oval(slide, l, t, 0.42, 0.42, color)
    tb(slide, l, t + 0.05, 0.42, 0.32, str(n), 14, True, WHITE, PP_ALIGN.CENTER)


def blank(prs, bg=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(bg_shape, bg)
    spTree = slide.shapes._spTree
    sp = bg_shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide

def build():
    cover = crop_cover(IMG / "cover-industry-water.jpg", CACHE / "cover.jpg")
    cert = crop_fit(IMG / "blue-certificate-generic.jpg", CACHE / "cert.jpg", (900, 680))
    team = crop_fit(IMG / "water-team.jpg", CACHE / "team.jpg", (900, 700))
    flow = crop_fit(IMG / "industry-water-flow.jpg", CACHE / "flow.jpg", (1400, 700))
    textile = crop_fit(IMG / "textile-plant.jpg", CACHE / "textile.jpg", (800, 620))
    food = crop_fit(IMG / "food-plant.jpg", CACHE / "food.jpg", (800, 620))
    metal = crop_fit(IMG / "metal-cooling.jpg", CACHE / "metal.jpg", (800, 620))
    rain = crop_fit(IMG / "rainwater.jpg", CACHE / "rain.jpg", (800, 620))
    small = crop_fit(IMG / "small-factory.jpg", CACHE / "small.jpg", (900, 700))
    plant = crop_fit(IMG / "cip-skid.jpg", CACHE / "plant.jpg", (900, 700))
    fixtures = crop_fit(IMG / "efficient-fixtures.jpg", CACHE / "fix.jpg", (800, 620))
    train = crop_fit(IMG / "training-session.jpg", CACHE / "train.jpg", (800, 620))
    posters = crop_fit(IMG / "awareness-posters.jpg", CACHE / "posters.jpg", (800, 620))

    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 KAPAK
    s = blank(prs, NAVY)
    picture(s, cover, 0, 0, 13.333, 7.5)
    rect(s, 0, 0, 13.333, 0.14, GOLD, rounded=False)
    pill(s, 0.55, 1.35, 5.35, 0.38, GOLD)
    tb(s, 0.55, 1.40, 5.35, 0.30, "SU VERİMLİLİĞİ  •  ENDÜSTRİYEL TESİSLER", 12, True, NAVY, PP_ALIGN.CENTER)
    tb(s, 0.55, 1.95, 8.2, 1.7, "Her tesis için\nMavi Su Belgesi", 36, True, WHITE)
    tb(s, 0.55, 3.75, 7.6, 1.1,
       "Tekstil, gıda, metal, kimya… NACE’niz ne olursa olsun:\nsu verimliliği sistemini kurmak, 7 kriteri sağlamak, belgeye başvurmak.",
       16, False, RGBColor(210, 228, 238))
    for i, (lab, col) in enumerate([("7 adım", TEAL), ("Örnekli anlatım", SKY), ("Zorunlu + gönüllü", GREEN)]):
        pill(s, 0.55 + i * 2.15, 5.15, 2.0, 0.38, col)
        tb(s, 0.55 + i * 2.15, 5.20, 2.0, 0.30, lab, 12, True, WHITE, PP_ALIGN.CENTER)
    tb(s, 0.55, 6.85, 8.2, 0.28, "Tarım ve Orman Bakanlığı  •  Su Verimliliği Yönetmeliği (RG 27.12.2024 / 32765)", 12, False, RGBColor(176, 200, 216))
    tb(s, 11.4, 6.85, 1.5, 0.28, "1 / 13", 12, True, RGBColor(176, 200, 216), PP_ALIGN.RIGHT)

    # 2 NEDEN
    s = blank(prs, CREAM)
    header_bar(s, "Bağlam", "Neden su verimliliği belgesi?",
               "Sanayi Türkiye su kullanımının ~%11’i. Belge, suyu ölçülebilir ve izlenebilir hale getirir.")
    picture(s, plant, 8.55, 1.55, 4.35, 3.35)
    rect(s, 8.55, 4.9, 4.35, 1.85, NAVY, radius=0.06)
    tb(s, 8.75, 5.05, 4.0, 0.35, "Ulusal hedef", 13, True, GOLD)
    tb(s, 8.75, 5.40, 4.0, 1.2,
       "Endüstride mevcut en iyi tekniklerle %50’ye varan su kazanımı hedefleniyor. 152 NACE için rehber var.",
       13, False, WHITE)
    stats = [
        ("%11", "Türkiye’de toplam suyun\nsanayideki payı", TEAL),
        ("%50’ye varan", "Sanayide ulusal su\nkazanımı hedefi", GOLD),
        ("~1.300 m³", "Kişi başı yıllık\nkullanılabilir su", SKY),
    ]
    for i, (num, desc, col) in enumerate(stats):
        card(s, 0.4 + i * 2.65, 1.55, 2.5, 2.35, WHITE)
        oval(s, 0.55 + i * 2.65, 1.72, 0.42, 0.42, col)
        tb(s, 0.55 + i * 2.65, 2.25, 2.2, 0.7, num, 18, True, NAVY)
        tb(s, 0.55 + i * 2.65, 2.95, 2.2, 0.75, desc, 12, False, MUTED)
    card(s, 0.4, 4.1, 7.9, 2.65, WHITE)
    tb(s, 0.6, 4.25, 7.5, 0.35, "Tipik bir fabrikada suyun gittiği yer (örnek dağılım)", 15, True, NAVY)
    bars = [
        ("Proses / yıkama", 0.32, TEAL, "%32"),
        ("Soğutma", 0.24, SKY, "%24"),
        ("Temizlik / CIP", 0.18, GOLD, "%18"),
        ("Buhar / kazan", 0.16, PURPLE, "%16"),
        ("Sosyal + peyzaj", 0.10, CORAL, "%10"),
    ]
    for i, (name, frac, col, lab) in enumerate(bars):
        y = 4.7 + i * 0.38
        tb(s, 0.6, y, 3.15, 0.32, name, 12, False, INK)
        rect(s, 3.8, y + 0.06, 3.6 * frac, 0.20, col, rounded=True, radius=0.4)
        tb(s, 7.45, y, 0.6, 0.32, lab, 12, True, col)
    footer(s, 2)

    # 3 SEVİYELER
    s = blank(prs, CREAM)
    header_bar(s, "Belge mimarisi", "Üç seviye: Mavi, Yeşil, Turkuaz",
               "Mavi = sistemi kurdunuz. Yeşil ve turkuaz = ölçülebilir kazanım. Tüm endüstriyel NACE’ler için aynı merdiven.")
    levels = [
        ("MAVİ", TEAL, "Sistem belgesi",
         "7 kriterle su verimliliği sistemini kurun.\nYükümlüler için ilk zorunlu basamak.\nGönüllüler de aynı 7 adımı izler.",
         "Geçerlilik: 5 yıl"),
        ("YEŞİL", GREEN, "Performans belgesi",
         "Kendi NACE rehberinizdeki teknikleri uygulayın.\nSuyun en az %10’u alternatif kaynaktan.\nISO 46001 ilk başvuruda aranmaz.",
         "Yükümlü: en fazla 5 yıl"),
        ("TURKUAZ", SKY, "İleri seviye (gönüllü)",
         "Atıksuyun en az %20’sini geri kazanın.\nRehber teknikleri + ISO 46001.\nISO 14046 su ayak izi belgesi.",
         "Tamamen gönüllü"),
    ]
    for i, (name, col, tag, body, note) in enumerate(levels):
        x = 0.4 + i * 4.25
        card(s, x, 1.55, 4.05, 5.2, WHITE)
        rect(s, x, 1.55, 4.05, 1.15, col, radius=0.08)
        tb(s, x, 1.68, 4.05, 0.5, name, 26, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x, 2.18, 4.05, 0.35, tag, 13, False, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.25, 2.9, 3.55, 2.4, body, 14, False, INK)
        pill(s, x + 0.55, 5.85, 2.95, 0.42, col)
        tb(s, x + 0.55, 5.91, 2.95, 0.32, note, 12, True, WHITE, PP_ALIGN.CENTER)
    footer(s, 3)

    # 4 YÜKÜMLÜ / GÖNÜLLÜ
    s = blank(prs, CREAM)
    header_bar(s, "Kapsam", "Tesisiniz yükümlü mü, gönüllü mü?",
               "Kural tüm sektörler için aynı: kılavuz Ek-1’deki NACE + çalışan sayısı. Liste: suverimliligi.gov.tr")
    card(s, 0.4, 1.55, 6.2, 5.2, WHITE)
    rect(s, 0.4, 1.55, 6.2, 0.7, CORAL, radius=0.08)
    tb(s, 0.4, 1.68, 6.2, 0.45, "YÜKÜMLÜ  •  50 ve üzeri çalışan", 18, True, WHITE, PP_ALIGN.CENTER)
    tb_multi(s, 0.65, 2.45, 5.75, 4.0, [
        ("Kim?", 13, True, CORAL, 0),
        ("Ek-1’de YÜKÜMLÜ yazan NACE’de faaliyet gösteren ve SGK’da 49’dan fazla çalışanı olan her endüstriyel tesis (tekstil, gıda, metal, kimya…).", 13, False, INK, 4),
        ("Ne yapılmalı?", 13, True, CORAL, 10),
        ("Yönetmelik yayımı itibarıyla 18 ay içinde sistemi kurup Mavi Belge’ye başvurmak.", 13, False, INK, 4),
        ("Sonra?", 13, True, CORAL, 10),
        ("En fazla 5 yıl içinde Yeşil Belge (Bakanlık onayıyla 7 yıla uzayabilir). İş termin planı Mavi dosyasında sunulur.", 13, False, INK, 4),
        ("Örnek: 180 kişilik boyahane veya dökümhane → zorunlu.", 13, True, NAVY, 12),
    ])
    card(s, 6.85, 1.55, 6.05, 5.2, WHITE)
    rect(s, 6.85, 1.55, 6.05, 0.7, GREEN, radius=0.08)
    tb(s, 6.85, 1.68, 6.05, 0.45, "GÖNÜLLÜ  •  50’nin altı veya istek", 18, True, WHITE, PP_ALIGN.CENTER)
    tb_multi(s, 7.1, 2.45, 5.55, 4.0, [
        ("Kim?", 13, True, GREEN, 0),
        ("Aynı NACE’de 50’den az çalışanı olan KOBİ’ler; Ek-1’de GÖNÜLLÜ yazan faaliyetler. Çalışan şartı aranmaz.", 13, False, INK, 4),
        ("Ne yapılmalı?", 13, True, GREEN, 10),
        ("Aynı 7 kriterle sistemi kurup csys.tarimorman.gov.tr üzerinden başvurabilirsiniz. Sıra bekleme yok.", 13, False, INK, 4),
        ("Avantaj", 13, True, GREEN, 10),
        ("Gönüllüler Yeşil veya Turkuaz’a doğrudan gidebilir. Müşteri, ihracat ve ihale dosyasında artı puandır.", 13, False, INK, 4),
        ("Örnek: 22 kişilik atölye → gönüllü, aynı CSYS formu.", 13, True, NAVY, 12),
    ])
    footer(s, 4)

    # 5 YEDI ADIM
    s = blank(prs, CREAM)
    header_bar(s, "Yol haritası", "Mavi Belge’nin 7 adımı",
               "Ek-3 kriterleri her NACE için aynıdır. Kanıt: atama yazısı, tablo, plan, imza listesi, fotoğraf.")
    steps_data = [
        ("1", "Ekip", "Lider + eğitim sorumlusu + teknik yardımcı", TEAL),
        ("2", "Mevcut durum", "Su-atıksu haritası, spesifik tüketim", SKY),
        ("3", "Hedefler", "5 yıllık, ölçülebilir, yeşile götüren", GOLD),
        ("4", "Plan", "İş termin planı, maliyet, yıl yıl eylem", AMBER),
        ("5", "Eğitim", "Bireysel + sektörel, yılda en az 1’er", PURPLE),
        ("6", "Ekipman", "Musluk, rezervuar, perlatör, sensör", GREEN),
        ("7", "Materyal", "Afiş, ekran, broşür, saha fotoğrafı", CORAL),
    ]
    for i, (n, title, desc, col) in enumerate(steps_data):
        y = 1.55 + i * 0.72
        card(s, 0.4, y, 8.35, 0.66, WHITE)
        oval(s, 0.55, y + 0.12, 0.42, 0.42, col)
        tb(s, 0.55, y + 0.18, 0.42, 0.32, n, 14, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 1.15, y + 0.08, 2.3, 0.48, title, 16, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, 3.5, y + 0.08, 5.0, 0.48, desc, 13, False, MUTED, anchor=MSO_ANCHOR.MIDDLE)
    picture(s, team, 8.95, 1.55, 3.95, 5.15)
    footer(s, 5)

    # 6 ADIM 1-2
    s = blank(prs, CREAM)
    header_bar(s, "Adım 1 ve 2", "Ekibi kurun, suyu görünür kılın",
               "Ölçülmeyen su yönetilemez. Sayaçlar proses, yardımcı ünite ve sosyal alanı ayrı konuşmalıdır.")
    card(s, 0.4, 1.5, 6.35, 2.55, WHITE)
    tb(s, 0.6, 1.62, 6.0, 0.32, "Adım 1  •  Su verimliliği ekibi", 16, True, TEAL)
    tb(s, 0.6, 2.0, 6.0, 1.85,
       "Kılavuz: 1 ekip lideri + en az 1 eğitim sorumlusu + en az 1 yardımcı.\n"
       "Bir kişi en fazla 2 rol alabilir. Lider dışındakiler yalnızca görüntüleyebilir.\n"
       "Örnek — Anadolu Tekstil: lider = işletme müdürü; eğitim = İK;\n"
       "yardımcı = bakım + AAT. Sistem e-posta ile aktivasyon yollar.",
       13, False, INK)
    card(s, 6.95, 1.5, 5.95, 2.55, WHITE)
    tb(s, 7.15, 1.62, 5.55, 0.32, "Adım 2  •  Mevcut durum", 16, True, SKY)
    tb(s, 7.15, 2.0, 5.55, 1.85,
       "CSYS’de 7 sekme: su (kaynak-kullanım-ücret), GOSK,\n"
       "peyzaj, eğitim, ekipman, materyal, sürdürülebilirlik.\n"
       "Örnek kaynak: kuyu 420.000 + şebeke 280.000 m³.\n"
       "Kütle denkliği: çekilen ≈ kullanılan + kayıp + ürün.",
       13, False, INK)
    tb(s, 0.45, 4.2, 12.4, 0.3, "Örnek mevcut durum — 180 kişilik boyahane, referans yıl", 14, True, NAVY)
    rows = [
        ("Proses / yıkama", "224.000 m³", "%32", GOLD),
        ("Soğutma", "168.000 m³", "%24", TEAL),
        ("Temizlik / CIP", "126.000 m³", "%18", SKY),
        ("Buhar / kazan", "112.000 m³", "%16", PURPLE),
        ("Sosyal + diğer", "70.000 m³", "%10", CORAL),
        ("TOPLAM", "700.000 m³", "100%", NAVY),
    ]
    for i, (a, b, c, col) in enumerate(rows):
        x = 0.4 + i * 2.15
        rect(s, x, 4.55, 2.05, 2.15, WHITE if i < 5 else NAVY, radius=0.08)
        tb(s, x + 0.08, 4.68, 1.9, 0.55, a, 11, True, col if i < 5 else GOLD, PP_ALIGN.CENTER)
        tb(s, x + 0.08, 5.25, 1.9, 0.5, b, 14, True, WHITE if i == 5 else NAVY, PP_ALIGN.CENTER)
        tb(s, x + 0.08, 5.8, 1.9, 0.55, c, 12, False, RGBColor(200, 214, 226) if i == 5 else MUTED, PP_ALIGN.CENTER)
    footer(s, 6)

    # 7 ADIM 3-4
    s = blank(prs, CREAM)
    header_bar(s, "Adım 3 ve 4", "Hedefi sayıya, planı takvime bağlayın",
               "CSYS Hedef Sepeti’nden seçin veya “Diğer” yazın. Her hedefe en az 1 eylem; termin en fazla +5 yıl.")
    card(s, 0.4, 1.5, 6.35, 5.25, WHITE)
    tb(s, 0.6, 1.65, 6.0, 0.32, "Örnek — boyahane hedef kaydı", 15, True, GOLD)
    goals = [
        ("MET ile toplam suyu azalt", "700.000 → 520.000 m³", "Sepet hedefi; gösterge m³/yıl"),
        ("Yağmur suyu oranını artır", "%0 → %4", "Termin: başvuru + 3 yıl"),
        ("Atıksu geri kazanımını artır", "%0 → %6", "AAT çıkışı kasa / zemin yıkamada"),
        ("Su kaybını düşür", "%9 → %5", "Kılavuz örneği: çekim − kullanım"),
    ]
    for i, (k, v, n) in enumerate(goals):
        y = 2.15 + i * 1.05
        oval(s, 0.65, y + 0.1, 0.18, 0.18, GOLD)
        tb(s, 1.0, y, 5.5, 0.28, k, 13, True, NAVY)
        tb(s, 1.0, y + 0.28, 5.5, 0.28, v, 16, True, TEAL)
        tb(s, 1.0, y + 0.56, 5.5, 0.28, n, 12, False, MUTED)
    card(s, 6.95, 1.5, 5.95, 5.25, WHITE)
    tb(s, 7.15, 1.65, 5.55, 0.32, "Her hedefe bağlı eylem (zorunlu)", 15, True, SKY)
    plan = [
        ("Yıl 1", "Sayaç ağı + kaçak turu + eğitimler", "Maliyet: 0,9 M TL"),
        ("Yıl 2", "Kapalı çevrim soğutma dönüşümü", "Maliyet: 3,2 M TL"),
        ("Yıl 3", "Çatı yağmur hasadı 250 m³ depo", "Maliyet: 1,8 M TL"),
        ("Yıl 4", "AAT çıkışını uygun proseslere bağla", "Maliyet: 2,6 M TL"),
        ("Yıl 5", "Yeşil belge dosyası + ISO 46001", "Maliyet: 0,4 M TL"),
    ]
    for i, (yr, act, save) in enumerate(plan):
        y = 2.12 + i * 0.85
        pill(s, 7.15, y, 1.05, 0.32, TEAL if i < 4 else GOLD)
        tb(s, 7.15, y + 0.03, 1.05, 0.26, yr, 11, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 8.3, y, 3.35, 0.55, act, 12, False, INK)
        tb(s, 8.3, y + 0.42, 3.35, 0.28, save, 11, True, MUTED)
    footer(s, 7)

    # 8 ADIM 5-7
    s = blank(prs, CREAM)
    header_bar(s, "Adım 5, 6 ve 7", "İnsan, armatür ve görünür hatırlatıcılar",
               "Mavi belge “sadece makine” değildir. Eğitim, musluk ve afiş her sektörde dosyanın parçasıdır.")
    blocks = [
        (train, TEAL, "5  Eğitim",
         "Başvurudan önce en az 1 bireysel + 1 sektörel eğitim.\n"
         "Bireysel: suverimliligi.gov.tr videosu; tüm personel 5 yılda.\n"
         "Sektörel: kendi NACE rehberiniz; lisans + 3 yıl deneyimli eğitmen.\n"
         "Kanıt: imza listesi, tutanak, foto — 5 yıl saklanır."),
        (fixtures, GREEN, "6  Su verimli ekipman",
         "Kapsam bireysel kullanım: musluk, duş, rezervuar, perlatör, sensör.\n"
         "Örnek: 42 lavaboya 6 L/dk perlatör, 18 tuvalete 6/3 L rezervuar.\n"
         "Sosyal su ~%30 düşer. Değişim 5 yıllık planda yazılır."),
        (posters, GOLD, "7  Yazılı-görsel materyal",
         "Afiş, broşür, kantin ekranı, yıkama odası talimatı.\n"
         "Örnek: “Hortumu açık bırakma — 1 dk = 15 L”.\n"
         "Kanıt: basılı örnek + saha fotoğrafı."),
    ]
    for i, (img, col, title, body) in enumerate(blocks):
        x = 0.35 + i * 4.3
        card(s, x, 1.5, 4.15, 5.25, WHITE)
        picture(s, img, x + 0.15, 1.65, 3.85, 1.85)
        tb(s, x + 0.2, 3.6, 3.75, 0.35, title, 16, True, col)
        tb(s, x + 0.2, 4.0, 3.75, 2.5, body, 12, False, INK)
    footer(s, 8)

    # 9 SAHA TEKNİKLERİ
    s = blank(prs, CREAM)
    header_bar(s, "Saha", "Her fabrikada suyu nerede kurtarırsınız?",
               "Önce kendi NACE rehberiniz (suverimliligi.gov.tr). Aşağıdakiler çoğu tesiste ortak MET’lerdir.")
    picture(s, flow, 0.4, 1.5, 6.4, 3.2)
    techniques = [
        ("Sayaç ve kaçak", "Proses / soğutma / buhar / sosyal ayrı sayılsın. Aylık sızıntı turu."),
        ("Kapalı soğutma", "Açık kule yerine kapalı çevrim; blowdown’u uygun yerde kullanın."),
        ("Yıkama / CIP", "Darbe durulama, iletkenlik kesimi, son durulama = sonraki ilk durulama."),
        ("Buhar kazanı", "Kondens dönüşü, blöf optimizasyonu, kaçak buhar avı."),
        ("Isı eşanjörü", "Isınan soğutma suyunu proses ön ısıtmasında depolayın."),
        ("Yağmur / AAT", "Çatı hasadı + arıtılmış su: zemin, kasa, yeşil alan, soğutma makyajı."),
    ]
    for i, (t, d) in enumerate(techniques):
        coln, row = i % 2, i // 2
        x = 6.95 + coln * 3.1
        y = 1.5 + row * 1.75
        card(s, x, y, 2.95, 1.62, WHITE)
        tb(s, x + 0.15, y + 0.12, 2.65, 0.4, t, 14, True, TEAL)
        tb(s, x + 0.15, y + 0.52, 2.65, 0.95, d, 12, False, INK)
    tb(s, 0.45, 4.85, 6.3, 0.3, "Temiz suyu kirli hatta karıştırmayın", 13, True, NAVY)
    tb(s, 0.45, 5.2, 6.3, 1.4,
       "İçme suyu kalitesindeki şebekeyi yer yıkamada kullanmayın. Geri kullanımın proses uygunluğunu kalite / gıda / ISO sorumlusu onaylasın. NACE rehberindeki öncelikli teknikler iş termin planına yazılır.",
       13, False, MUTED)
    footer(s, 9)

    # 10 ÜÇ VAKA
    s = blank(prs, CREAM)
    header_bar(s, "Üç sektör, aynı mantık", "Tekstil, gıda, metal — üç kazanç hikâyesi",
               "Rakamlar öğretici senaryodur; kendi sayaç verinizle yeniden hesaplayın.")
    cases = [
        (textile, TEAL, "Tekstil boyahane",
         "Durum: sürekli durulama, iletkenlik yok.\n"
         "İş: darbe durulama + son durulama tankı.\n"
         "Sonuç: durulama suyu ~%40, tesis toplamı ~%10.\n"
         "Belgeye: SOP, önce-sonra sayaç, eğitim tutanağı."),
        (food, GOLD, "Gıda işleme",
         "Durum: pastörizatör / yıkama suyu gideri.\n"
         "İş: taşma suyu geri dönüş + hortum tabancası.\n"
         "Sonuç: hat bazında %15–25 taze su düşer.\n"
         "Belgeye: P&ID, fotoğraf, aylık m³."),
        (metal, SKY, "Metal / soğutma",
         "Durum: açık kule, yüksek blöf.\n"
         "İş: kapalı çevrim + blöfün zemin yıkamada kullanımı.\n"
         "Sonuç: makyaj suyu ciddi azalır; yağmur ile GOSK payı artar.\n"
         "Belgeye: debi, analiz, kullanım noktası listesi."),
    ]
    for i, (img, col, title, body) in enumerate(cases):
        x = 0.35 + i * 4.3
        card(s, x, 1.5, 4.15, 5.25, WHITE)
        picture(s, img, x + 0.12, 1.62, 3.91, 1.7)
        rect(s, x + 0.12, 3.32, 3.91, 0.08, col, rounded=False)
        tb(s, x + 0.2, 3.48, 3.75, 0.4, title, 16, True, col)
        tb(s, x + 0.2, 3.92, 3.75, 2.55, body, 12, False, INK)
    footer(s, 10)

    # 11 İKİ SENARYO
    s = blank(prs, CREAM)
    header_bar(s, "İki tesis, aynı 7 adım", "Büyük fabrika ve küçük atölye nasıl ilerler?",
               "Ölçek değişir, kriter değişmez. Dosya tesis bazındadır: ayrı lokasyon = ayrı başvuru.")
    picture(s, textile, 0.4, 1.5, 4.15, 2.45)
    picture(s, small, 0.4, 4.15, 4.15, 2.55)
    card(s, 4.75, 1.5, 8.15, 2.45, WHITE)
    tb(s, 4.95, 1.62, 7.8, 0.32, "Anadolu Tekstil A.Ş.  •  180 çalışan  •  YÜKÜMLÜ", 15, True, CORAL)
    tb(s, 4.95, 2.05, 7.8, 1.7,
       "18 ay içinde sistemi kurup önce yetki, sonra Mavi belge başvurusu yapar. "
       "CSYS’ye 4’lü/6’lı NACE, ürün, 6 zorunlu evrak ve 5 yıllık hedef-eylem girer. "
       "OSB içinde olsa bile kendi belgesini alır. İki fabrika varsa iki yetki, iki belge.",
       13, False, INK)
    card(s, 4.75, 4.15, 8.15, 2.55, WHITE)
    tb(s, 4.95, 4.27, 7.8, 0.32, "Mavi Metal Atölye  •  22 çalışan  •  GÖNÜLLÜ", 15, True, GREEN)
    tb(s, 4.95, 4.7, 7.8, 1.8,
       "Zorunluluk yok; aynı CSYS formu. 3 kişilik ekip, 1 proses sayacı + fatura, "
       "hedef sepetinden “MET ile suyu azalt”, hortum tabancası ve perlatör fotoğrafı. "
       "Yetki onayı yine şart. Yeşile atlamak isterse sıralı beklemez.",
       13, False, INK)
    footer(s, 11)

    # 12 BAŞVURU
    s = blank(prs, CREAM)
    header_bar(s, "Dijital süreç", "Önce yetki, sonra belge — yalnızca CSYS",
               "Kılavuz 3.1.1.1  •  csys.tarimorman.gov.tr  •  e-Devlet vatandaş girişi  •  kargo/e-posta kabul edilmez")
    picture(s, cert, 9.2, 1.48, 3.7, 2.55)
    card(s, 0.4, 1.48, 8.55, 2.55, WHITE)
    pill(s, 0.6, 1.62, 1.55, 0.32, GOLD)
    tb(s, 0.6, 1.66, 1.55, 0.26, "A. YETKİ", 12, True, NAVY, PP_ALIGN.CENTER)
    tb(s, 2.25, 1.62, 6.5, 0.7,
       "Üç unsur bağlanır: tüzel kişilik + fiziki tesis/fabrika + yetkili kişi. Onay e-postası gelmeden belge ekranı açılmaz.",
       13, False, INK)
    pill(s, 0.6, 2.45, 1.55, 0.32, TEAL)
    tb(s, 0.6, 2.49, 1.55, 0.26, "B. BELGE", 12, True, WHITE, PP_ALIGN.CENTER)
    tb(s, 2.25, 2.42, 6.5, 1.4,
       "5 bölüm sırayla yeşile döner: 1) başvuru sahibi (genel + 4’lü/6’lı NACE üretim + evrak)  2) ekip  3) mevcut durum  4) hedefler-eylemler  5) PDF kabul + Başvuruyu Gönder. Kaydet deyip çıkabilirsiniz.",
       13, False, INK)
    docs = [
        ("1", "SGK çalışan sayısı", "Başvuru yılı tahakkuk / barkodlu nüsha"),
        ("2", "Oda faaliyet belgesi", "Son 1 yıl, unvan-adres-üretim"),
        ("3", "Kapasite raporu", "Geçerli; beyan edilen NACE yazmalı"),
        ("4", "Sanayi sicil belgesi", "Güncel; unvan-adres-faaliyet"),
        ("5", "Ticaret sicil gazetesi", "Şube ise şube bilgisi şart"),
        ("6", "Çevre izin / muafiyet", "Güncel izin veya resmi muafiyet yazısı"),
    ]
    for i, (n, t, d) in enumerate(docs):
        col, row = i % 3, i // 3
        x = 0.4 + col * 4.25
        y = 4.2 + row * 1.35
        card(s, x, y, 4.1, 1.22, WHITE)
        oval(s, x + 0.14, y + 0.36, 0.42, 0.42, TEAL)
        tb(s, x + 0.14, y + 0.43, 0.42, 0.3, n, 14, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.68, y + 0.18, 3.25, 0.4, t, 14, True, NAVY)
        tb(s, x + 0.68, y + 0.58, 3.25, 0.5, d, 12, False, MUTED)
    footer(s, 12)

    # 13 KONTROL
    s = blank(prs, CREAM)
    header_bar(s, "Kapanış", "Başvuru öncesi 12 maddelik kontrol listesi",
               "Hepsi “evet” ise CSYS dosyanız olgun demektir. Dosya ≤15 MB; pdf / jpg / png / tiff.")
    checks = [
        ("NACE (4’lü + 6’lı) fiili üretim ve kapasite raporuyla uyumlu.", TEAL),
        ("SGK: başvuru yılı. 49’dan fazla çalışan + YÜKÜMLÜ NACE = zorunlu.", SKY),
        ("Yetki onayı alınmış: tüzel kişi + tesis + yetkili eşleşiyor.", GOLD),
        ("Ekip: 1 lider + ≥1 eğitim + ≥1 yardımcı; kişi başı en fazla 2 rol.", PURPLE),
        ("Mevcut durum 7 sekmesi ve su kütle denkliği dolu.", GREEN),
        ("Hedef sepeti + her hedefe ≥1 eylem + 5 yıl termin + maliyet.", CORAL),
        ("En az 1 bireysel + 1 sektörel eğitim; imza listesi yüklü.", TEAL),
        ("Musluk-rezervuar-perlatör ve afiş fotoğrafları sistemde.", SKY),
        ("6 zorunlu evrak: SGK, oda, kapasite, sicil, gazete, çevre izin.", GOLD),
        ("Ayrı fabrikalar için ayrı yetki / ayrı belge düşünüldü.", PURPLE),
        ("PDF önizleme kabul edildi; “verilerin doğruluğunu taahhüt”.", GREEN),
        ("Statü: incelemede geri alınmaz. 2. revize = red; 60 gün kuralı.", CORAL),
    ]
    for i, (txt, col) in enumerate(checks):
        coln, row = i % 2, i // 2
        x = 0.4 + coln * 6.45
        y = 1.48 + row * 0.85
        card(s, x, y, 6.25, 0.75, WHITE)
        oval(s, x + 0.15, y + 0.18, 0.38, 0.38, col)
        tb(s, x + 0.15, y + 0.24, 0.38, 0.28, "✓", 14, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.65, y + 0.12, 5.4, 0.52, txt, 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 0.4, 7.16, 11.2, 0.28,
       "Kaynak: Mavi Belge Başvuru Kılavuzu 3.1.1.1  •  Su Verimliliği Yönetmeliği (RG 27.12.2024/32765)  •  suverimliligi.gov.tr  •  csys.tarimorman.gov.tr",
       10, False, MUTED)
    tb(s, 12.2, 7.16, 0.8, 0.25, "13/13", 10, True, MUTED, PP_ALIGN.RIGHT)
    prs.save(OUT)
    print(f"Saved {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    build()
