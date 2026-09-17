#!/usr/bin/env python3
"""KKDİK Yönetmeliği bilgilendirme sunumu (kimler kapsar, süreç, ne yapılmalı)."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "KKDIK_Yonetmeligi_Rehberi.pptx"

W, H = Inches(13.333), Inches(7.5)

NAVY = RGBColor(12, 32, 64)
NAVY2 = RGBColor(18, 48, 92)
TEAL = RGBColor(0, 140, 142)
TEAL_D = RGBColor(0, 108, 112)
GOLD = RGBColor(201, 148, 42)
AMBER = RGBColor(196, 112, 28)
CREAM = RGBColor(246, 242, 232)
WHITE = RGBColor(255, 255, 255)
SOFT = RGBColor(232, 244, 247)
MINT = RGBColor(228, 245, 236)
SAND = RGBColor(255, 244, 224)
LAV = RGBColor(236, 232, 252)
INK = RGBColor(28, 39, 54)
MUTED = RGBColor(90, 104, 122)
GREEN = RGBColor(26, 138, 99)
CORAL = RGBColor(196, 64, 64)
SKY = RGBColor(36, 120, 196)
PURPLE = RGBColor(92, 78, 168)
FONT = "Calibri"


def set_run(run, size, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._r.get_or_add_rPr()
    latin = r_pr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(r_pr, qn("a:latin"))
    latin.set("typeface", FONT)


def add_text(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, italic)
    return p


def add_para(tf, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, space_before=0, space_after=4):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return p


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, left, top, width, height, color, rounded=True, radius=0.08):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    fill(shp, color)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def pill(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    fill(shp, color)
    try:
        shp.adjustments[0] = 0.5
    except Exception:
        pass
    return shp


def oval(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    fill(shp, color)
    return shp


def tb(slide, left, top, width, height, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf._txBody.bodyPr.set(
        "anchor",
        {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"),
    )
    add_text(tf, text, size, bold, color, align, italic)
    return box


def blank(prs, color=CREAM):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(bg, color)
    return slide


def header_bar(slide, kicker, title, subtitle=""):
    rect(slide, 0, 0, 13.333, 1.22, NAVY, rounded=False)
    rect(slide, 0, 1.22, 13.333, 0.08, TEAL, rounded=False)
    pill(slide, 0.4, 0.16, 2.35, 0.28, TEAL)
    tb(slide, 0.4, 0.18, 2.35, 0.24, kicker, 11, True, WHITE, PP_ALIGN.CENTER)
    tb(slide, 0.4, 0.48, 12.4, 0.42, title, 24, True, WHITE)
    if subtitle:
        tb(slide, 0.4, 0.9, 12.4, 0.26, subtitle, 12, False, RGBColor(190, 214, 230))


def footer(slide, page, total=12):
    tb(
        slide,
        0.4,
        7.16,
        11.2,
        0.26,
        "KKDİK  •  RG 23.06.2017/30105 Mükerrer  •  Değişiklik RG 23.12.2023/32408  •  Usul ve Esaslar 05.08.2025  •  kimyasallar.csb.gov.tr",
        10,
        False,
        MUTED,
    )
    tb(slide, 12.05, 7.16, 0.9, 0.25, f"{page}/{total}", 10, True, MUTED, PP_ALIGN.RIGHT)


def card(slide, left, top, width, height, color=WHITE):
    return rect(slide, left, top, width, height, color, rounded=True, radius=0.06)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 KAPAK
    s = blank(prs, NAVY)
    rect(s, 0, 0, 0.22, 7.5, TEAL, rounded=False)
    rect(s, 10.9, 0, 2.44, 7.5, NAVY2, rounded=False)
    pill(s, 0.55, 1.15, 3.55, 0.38, TEAL)
    tb(s, 0.55, 1.2, 3.55, 0.3, "ÇŞİDB  •  KİMYASAL KAYIT SİSTEMİ", 12, True, WHITE, PP_ALIGN.CENTER)
    tb(s, 0.55, 1.75, 10.0, 1.15, "KKDİK Yönetmeliği", 44, True, WHITE)
    tb(
        s,
        0.55,
        2.9,
        10.0,
        0.45,
        "Kimyasalların Kaydı, Değerlendirilmesi, İzni ve Kısıtlanması",
        20,
        False,
        RGBColor(186, 214, 230),
    )
    tb(
        s,
        0.55,
        3.45,
        9.8,
        0.7,
        "Kimler kapsama girer, sistem nasıl ilerler ve kapsamdaysanız bugün ne yapmalısınız.",
        18,
        False,
        WHITE,
    )
    pill(s, 0.55, 4.4, 5.55, 0.5, GOLD)
    tb(s, 0.55, 4.48, 5.55, 0.36, "Kritik eşik  •  30 Eylül 2026 geçici kayıt", 15, True, NAVY, PP_ALIGN.CENTER)
    points = [
        ("01", "Kapsam"),
        ("02", "Süreç"),
        ("03", "Ne yapmalı"),
    ]
    for i, (num, label) in enumerate(points):
        y = 1.5 + i * 1.55
        oval(s, 11.55, y, 0.85, 0.85, TEAL)
        tb(s, 11.55, y + 0.22, 0.85, 0.42, num, 18, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 11.35, y + 0.92, 1.25, 0.35, label, 13, True, WHITE, PP_ALIGN.CENTER)
    tb(s, 0.55, 6.85, 9.5, 0.3, "Eylül 2026  •  Genel bilgilendirme  •  Resmî başvuru KKS üzerinden yapılır", 12, False, RGBColor(160, 186, 204))

    # 2 NEDİR
    s = blank(prs)
    header_bar(
        s,
        "Temel",
        "KKDİK nedir, neyi düzenler?",
        "Türkiye’nin kimyasal güvenlik sistemi  •  AB REACH ile aynı mantık, ayrı ulusal kayıt",
    )
    pillars = [
        ("Kayıt", "K", "Yılda 1 ton ve üzeri maddeyi imal veya ithal eden, teknik dosyayı KKS’ye verir. Veri yoksa madde piyasada duramaz."),
        ("Değerlendirme", "D", "Bakanlık dosyayı inceler; eksik veri, test ve kimyasal güvenlik raporunu isteyebilir."),
        ("İzin", "İ", "Çok yüksek önem arz eden maddeler (SVHC / Ek-14) izne bağlanır; izinsiz kullanım kısıtlanır."),
        ("Kısıtlama", "K", "Ek-17’deki maddelerin imalatı, piyasaya arzı veya kullanımı yasaklanır ya da şartlara bağlanır."),
    ]
    for i, (title, letter, body) in enumerate(pillars):
        x = 0.4 + i * 3.2
        card(s, x, 1.55, 3.05, 3.55, WHITE)
        oval(s, x + 1.1, 1.75, 0.85, 0.85, TEAL if i % 2 == 0 else NAVY)
        tb(s, x + 1.1, 1.95, 0.85, 0.5, letter, 22, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.18, 2.75, 2.7, 0.4, title, 18, True, NAVY, PP_ALIGN.CENTER)
        tb(s, x + 0.18, 3.25, 2.7, 1.65, body, 13, False, INK)
    card(s, 0.4, 5.25, 12.5, 1.7, SAND)
    tb(s, 0.65, 5.4, 12.1, 0.35, "Pratik kural", 14, True, AMBER)
    tb(
        s,
        0.65,
        5.8,
        12.1,
        0.95,
        "KKDİK “tehlikeli madde yönetmeliği” değildir. Tehlikesiz görünen bir hammadde de, muaf değilse ve yılda 1 tonu aşıyorsa kayıt ister. "
        "Kayıt maddeye değil firmaya aittir: tedarikçinizin kaydı, sizin ithalat kaydınızın yerine geçmez — tek istisna, yabancı üreticinin atadığı tek temsilcidir.",
        14,
        False,
        INK,
    )
    footer(s, 2)

    # 3 KİMLER
    s = blank(prs)
    header_bar(
        s,
        "Kapsam",
        "Kimler kapsama girer? Rolünüz kaydı belirler",
        "Kayıt ettiren yalnızca Türkiye’de yerleşik gerçek veya tüzel kişidir",
    )
    roles = [
        ("İmalatçı", TEAL, "Türkiye’de maddeyi kendi halinde üretir. Yılda ≥1 ton ise kayıttan sorumludur."),
        ("İthalatçı", SKY, "Maddeyi veya karışımı gümrükten kendi adına sokar. Tek temsilci yoksa kayıt ithalatçının."),
        ("Tek temsilci", GOLD, "Yabancı üretici, Türkiye’de bir kişiyi ithalatçı yükümlülüğü için atar. Atama sizi kapsıyorsa ithalatçı alt kullanıcı olur."),
        ("Alt kullanıcı", GREEN, "Maddeyi prosesinde kullanır; yurt içi tedarikçiden alır. Kayıt yapmaz, GBF ve kullanım koşullarına uyar."),
        ("Dağıtıcı", PURPLE, "Depolar ve piyasaya arz eder, kendisi kullanmaz. Kayıt yok; tedarik zinciri bilgisini iletir."),
        ("Eşya üretici / ithalatçı", CORAL, "Boru, parça, mamul. Kayıt yalnızca madde kasıtlı salınıyorsa ve ≥1 ton ise. SVHC ≥ %0,1 ise bildirim doğabilir."),
    ]
    for i, (title, color, body) in enumerate(roles):
        col, row = i % 3, i // 3
        x = 0.4 + col * 4.25
        y = 1.5 + row * 2.15
        card(s, x, y, 4.1, 2.0, WHITE)
        rect(s, x, y, 0.14, 2.0, color, rounded=False)
        tb(s, x + 0.35, y + 0.18, 3.55, 0.4, title, 16, True, NAVY)
        tb(s, x + 0.35, y + 0.62, 3.55, 1.2, body, 13, False, INK)
    footer(s, 3)

    # 4 5 ADIM
    s = blank(prs)
    header_bar(
        s,
        "Karar ağacı",
        "Bir madde kapsamda mı? Beş adımda bakın",
        "Karışım kayıt edilmez; içindeki her madde ayrı ayrı tartılır",
    )
    steps = [
        ("1", "Kimlik", "Ticari ad yetmez. GBF bölüm 3’ten CAS / EC alın. Hidrat, tuz ve farklı saflık ayrı madde olabilir."),
        ("2", "Muafiyet", "Atık, radyoaktif, transit, ilaç-gıda (kısmi), Ek-4 / Ek-5, polimerin kendisi ve <1 ton kayıt dışıdır."),
        ("3", "Tonaj", "Tüzel kişi + madde + takvim yılı. Üç tedarikçi 400’er kg ise toplam 1,2 tondur — eşik aşılır."),
        ("4", "Rol", "İmal / ithal / tek temsilci = kayıt. Yurt içi alım = alt kullanıcı. Tek temsilci sizi kapsıyorsa kayıt düşer."),
        ("5", "Son tarih", "Geçici kayıt: 30 Eylül 2026. Tam kayıt: tonaj ve CMR / sucul sınıfa göre 2026, 2028 veya 2030."),
    ]
    for i, (num, title, body) in enumerate(steps):
        y = 1.48 + i * 1.05
        card(s, 0.4, y, 12.5, 0.95, WHITE)
        oval(s, 0.58, y + 0.22, 0.52, 0.52, TEAL)
        tb(s, 0.58, y + 0.32, 0.52, 0.35, num, 16, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 1.3, y + 0.12, 2.2, 0.7, title, 18, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, 3.55, y + 0.16, 9.1, 0.65, body, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 4)

    # 5 MUAFİYET
    s = blank(prs)
    header_bar(
        s,
        "Sınır",
        "Kimler / neler kapsam dışında kalır?",
        "Muafiyet yazılı dayanak olmadan “kapsam dışı” denemez  •  GBF ve SEA çoğu muafiyette devam eder",
    )
    card(s, 0.4, 1.5, 6.2, 5.35, WHITE)
    pill(s, 0.6, 1.68, 2.7, 0.34, CORAL)
    tb(s, 0.6, 1.72, 2.7, 0.28, "Yönetmelik kapsamı dışı", 13, True, WHITE, PP_ALIGN.CENTER)
    left_items = [
        "Radyoaktif maddeler ve karışımlar",
        "İşlemsiz gümrük transiti / serbest bölge depolama",
        "İzole edilmemiş ara maddeler",
        "Taşıma (kara, demir, deniz, hava, içsu)",
        "Atık statüsündeki malzemeler",
        "Savunma amaçlı madde, karışım ve eşya",
    ]
    for i, item in enumerate(left_items):
        y = 2.2 + i * 0.7
        oval(s, 0.7, y + 0.08, 0.28, 0.28, CORAL)
        tb(s, 0.7, y + 0.1, 0.28, 0.24, str(i + 1), 11, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 1.15, y, 5.2, 0.5, item, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 6.8, 1.5, 6.1, 5.35, WHITE)
    pill(s, 7.0, 1.68, 3.35, 0.34, GREEN)
    tb(s, 7.0, 1.72, 3.35, 0.28, "Kayıttan (kısmen) muaf", 13, True, WHITE, PP_ALIGN.CENTER)
    right_items = [
        "Ek-4 listedeki maddeler (ör. bazı doğal maddeler)",
        "Ek-5 muafiyetleri (belirli doğal, yan ürün vb.)",
        "Polimerin kendisi (monomer ≥ %2 ve ≥1 ton ise kayıt)",
        "İlaç, veteriner, gıda-yem amaçlı kullanım (kısmi)",
        "PPORD (ürün ve süreç odaklı Ar-Ge) bildirimi",
        "Yılda 1 tonun altındaki maddeler (kayıt eşiği)",
    ]
    for i, item in enumerate(right_items):
        y = 2.2 + i * 0.7
        oval(s, 7.1, y + 0.08, 0.28, 0.28, GREEN)
        tb(s, 7.1, y + 0.1, 0.28, 0.24, str(i + 1), 11, True, WHITE, PP_ALIGN.CENTER)
        tb(s, 7.55, y, 5.1, 0.5, item, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 5)

    # 6 SİSTEM
    s = blank(prs)
    header_bar(
        s,
        "Süreç",
        "Sistem nasıl ilerler? KKS üzerinden yedi adım",
        "Tek madde — tek kayıt  •  Ortak kayıt esastır  •  Dosyayı Kimyasal Değerlendirme Uzmanı (KDU) girer",
    )
    flow = [
        ("1", "Envanter", "CAS bazında madde listesi, GBF, yıllık tonaj belgesi"),
        ("2", "EÇBS / KKS", "Entegre Çevre Bilgi Sistemi hesabı, yetkili kişi"),
        ("3", "Ön-MBDF", "Aynı maddeyi kaydedeceklerle forum; veri paylaşımı kapısı"),
        ("4", "Lider / üye", "Ortak kayıt grubu, sözleşme, LoA / maliyet paylaşımı"),
        ("5", "Geçici kayıt", "30.09.2026’ya kadar en az bir geçici kayıt numarası"),
        ("6", "Tam kayıt", "Tonaj ekleri (Ek-7…10), gerekirse KGR / CSR"),
        ("7", "GBF + izleme", "Türkçe 16 bölümlü GBF, güncelleme, izin / kısıtlama"),
    ]
    for i, (num, title, body) in enumerate(flow):
        x = 0.35 + i * 1.85
        card(s, x, 1.55, 1.75, 3.55, WHITE)
        oval(s, x + 0.52, 1.75, 0.7, 0.7, TEAL if i < 5 else GOLD)
        tb(s, x + 0.52, 1.9, 0.7, 0.42, num, 18, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.08, 2.6, 1.6, 0.7, title, 14, True, NAVY, PP_ALIGN.CENTER)
        tb(s, x + 0.08, 3.35, 1.6, 1.5, body, 12, False, INK, PP_ALIGN.CENTER)
        if i < 6:
            tb(s, x + 1.55, 2.85, 0.35, 0.3, "→", 16, True, TEAL, PP_ALIGN.CENTER)
    card(s, 0.4, 5.3, 12.5, 1.65, SOFT)
    tb(s, 0.65, 5.45, 12.1, 0.32, "KKS nerede çalışır?", 14, True, NAVY)
    tb(
        s,
        0.65,
        5.85,
        12.1,
        0.9,
        "Kayıt dosyası EÇBS içindeki Kimyasal Kayıt Sistemi’ne yüklenir. Geçici kayıt şablonu: Madde Yönetimi → Kayıt → KKDİK Geçici Kayıt. "
        "Gönderim sonrası döner sermaye referansı oluşur; ödeme KKS’de “Ödeme Yapıldı: Evet” görünmeden dosya değerlendirmeye alınmaz. "
        "İletişim: kimyasallar@csb.gov.tr  •  kimyasallar.csb.gov.tr",
        14,
        False,
        INK,
    )
    footer(s, 6)

    # 7 TAKVİM
    s = blank(prs)
    header_bar(
        s,
        "Takvim",
        "İki zincir: geçici kayıt 2026, tam kayıt 2026–2030",
        "En sık hata: tam kayıt 2030 diye 30 Eylül 2026 geçici kaydı atlamak",
    )
    card(s, 0.4, 1.5, 12.5, 1.55, SAND)
    pill(s, 0.6, 1.68, 3.4, 0.32, GOLD)
    tb(s, 0.6, 1.72, 3.4, 0.26, "Şu anki acil eşik", 12, True, NAVY, PP_ALIGN.CENTER)
    tb(s, 4.15, 1.65, 8.5, 0.4, "30 Eylül 2026  —  üye / bireysel geçici kayıt", 22, True, NAVY)
    tb(
        s,
        0.6,
        2.15,
        12.05,
        0.7,
        "Yılda ≥1 ton imal veya ithal eden her firma, tonajdan bağımsız, bu tarihe kadar en az bir geçici kayıt numarası almış olmalıdır. "
        "Lider geçici kayıt tarihi (31.03.2026) geçti; lider yoksa gerekçeli bireysel geçici kayıt KKS’den gönderilir. Numarasız madde bu tarihten sonra piyasaya arz edilemez.",
        13,
        False,
        INK,
    )
    rows = [
        ("31.12.2026", "≥1.000 t/yıl  •  CMR 1A/1B (≥1 t)  •  Sucul Akut 1 / Kronik 1 (≥100 t, H400/H410)", CORAL),
        ("31.12.2028", "≥100 t/yıl  (yukarıdaki özel zararlılık ölçütü taşımayan maddeler)", SKY),
        ("31.12.2030", "≥1 t/yıl  (1–100 ton bandındaki standart maddeler)", GREEN),
    ]
    tb(s, 0.4, 3.2, 12.5, 0.32, "Tam kayıt son tarihleri  (Yönetmelik geçici madde — RG 23.12.2023/32408)", 13, True, NAVY)
    for i, (date, desc, color) in enumerate(rows):
        y = 3.6 + i * 0.95
        card(s, 0.4, y, 12.5, 0.85, WHITE)
        rect(s, 0.4, y, 0.16, 0.85, color, rounded=False)
        tb(s, 0.75, y + 0.2, 2.4, 0.45, date, 20, True, color, anchor=MSO_ANCHOR.MIDDLE)
        tb(s, 3.3, y + 0.15, 9.3, 0.55, desc, 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 7)

    # 8 GEÇİCİ vs TAM
    s = blank(prs)
    header_bar(
        s,
        "Dosya",
        "Geçici kayıt ile tam kayıt aynı iş değil",
        "Usul ve Esaslar Ek-1  •  Geçici kayıtta da döner sermaye harcı alınır  •  KDU belgesi eklenir",
    )
    card(s, 0.4, 1.5, 6.2, 5.35, WHITE)
    pill(s, 0.6, 1.68, 2.6, 0.34, GOLD)
    tb(s, 0.6, 1.72, 2.6, 0.28, "Geçici kayıt", 13, True, NAVY, PP_ALIGN.CENTER)
    gecici = [
        "Amaç: piyasaya arzın kesilmemesi",
        "Son tarih: 30.09.2026 (üye / bireysel)",
        "KKS şablonu: KKDİK Geçici Kayıt",
        "Madde kimliği, bileşim, analitik bilgi",
        "Kullanım, tesis, tahmini tonaj, GHS",
        "Fizikokimya + güvenli kullanım özeti",
        "Ortak başvuruda değilseniz gerekçe yazın",
        "Ödeme referansı → KKS’de “ödeme evet”",
    ]
    for i, item in enumerate(gecici):
        tb(s, 0.7, 2.15 + i * 0.52, 5.7, 0.5, f"•  {item}", 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    card(s, 6.8, 1.5, 6.1, 5.35, WHITE)
    pill(s, 7.0, 1.68, 2.35, 0.34, TEAL)
    tb(s, 7.0, 1.72, 2.35, 0.28, "Tam kayıt", 13, True, WHITE, PP_ALIGN.CENTER)
    tam = [
        "Amaç: kalıcı kayıt numarası",
        "Son tarih: 2026 / 2028 / 2030 bandı",
        "Teknik dosya + tonaja göre Ek-7…10",
        "≥10 t/yıl: Kimyasal Güvenlik Raporu",
        "Çalışma özetleri, veri paylaşımı / LoA",
        "Lider dosyası + üye bilgisi (ortak kayıt)",
        "Veri yoksa gerekçeli ek süre (en fazla +2 yıl)",
        "Kayıt sonrası 3 ay içinde güncelleme kuralı",
    ]
    for i, item in enumerate(tam):
        tb(s, 7.1, 2.15 + i * 0.52, 5.55, 0.5, f"•  {item}", 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 8)

    # 9 İMALATÇI / İTHALATÇI
    s = blank(prs)
    header_bar(
        s,
        "Yükümlü",
        "İmalatçı, ithalatçı veya tek temsilciyseniz",
        "Kayıt sizin omzunuzdadır  •  30 Eylül 2026 öncesi envanter + KKS + geçici kayıt",
    )
    actions = [
        ("1", "Madde envanteri", "Tüm ithal / imal kimyasalları CAS ile listeleyin. Karışımda her bileşeni % ve tonaja çevirin."),
        ("2", "Rol teyidi", "Gümrük beyannamesinde ithalatçı siz misiniz? Tek temsilci atama yazısı sizi kapsıyor mu?"),
        ("3", "Muafiyet notu", "Kapsam dışı diyorsanız Ek madde numarasını dosyaya yazın. Yazılı dayanak yoksa kapsam varsayılır."),
        ("4", "KKS + ön-MBDF", "EÇBS hesabı, madde sorgu, aynı maddeyi kaydedeceklerle temas, ortak kayıt tercihi."),
        ("5", "Geçici kayıt", "Lider yoksa bireysel geçici kayıt + gerekçe. KDU girişi, harç, ödeme teyidi."),
        ("6", "Tam kayıt planı", "Tonaj bandı ve CMR/sucul sınıfa göre 2026-2030 dosya, veri ve bütçe takvimi."),
    ]
    for i, (num, title, body) in enumerate(actions):
        col, row = i % 2, i // 2
        x = 0.4 + col * 6.45
        y = 1.5 + row * 1.7
        card(s, x, y, 6.25, 1.55, WHITE)
        oval(s, x + 0.2, y + 0.5, 0.52, 0.52, TEAL)
        tb(s, x + 0.2, y + 0.6, 0.52, 0.35, num, 16, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.9, y + 0.18, 5.1, 0.4, title, 16, True, NAVY)
        tb(s, x + 0.9, y + 0.62, 5.1, 0.75, body, 13, False, INK)
    footer(s, 9)

    # 10 ALT KULLANICI / EŞYA
    s = blank(prs)
    header_bar(
        s,
        "Diğer roller",
        "Alt kullanıcı, dağıtıcı veya eşya üreticisiyseniz",
        "Kayıt yok diye yükümlülük yok demek değildir",
    )
    blocks = [
        (
            "Alt kullanıcı",
            GREEN,
            [
                "Yurt içi tedarikçiden alınan madde / karışım",
                "Kayıt tedarikçinin (veya onun ithalatçısının)",
                "Türkçe KKDİK formatlı GBF isteyin ve saklayın",
                "Öngörülen kullanıma uyun; sapmayı bildirin",
                "İzne / kısıtlamaya tabi madde var mı kontrol edin",
                "Tonajınız ithalata dönerse rol değişir — izleyin",
            ],
        ),
        (
            "Dağıtıcı",
            PURPLE,
            [
                "Kendiniz kullanmazsınız, depolar ve satarsınız",
                "GBF ve etiket bilgisini müşteriye aktarın",
                "Tedarik zincirindeki kayıt numarasını sorun",
                "Yeniden ambalaj / yeniden etiket SEA’ya uyar",
                "İthalata geçerseniz ithalatçı olursunuz",
                "Perakende de bu tanıma dahildir",
            ],
        ),
        (
            "Eşya üretici / ithalatçı",
            CORAL,
            [
                "Boru, panel, makine, ambalajlı mamul",
                "Kasıtlı salınım + ≥1 ton → madde kaydı",
                "SVHC ≥ %0,1 (a/a) ve ≥1 ton → bildirim",
                "CTP / plastik: monomer ve katkıları ayrı bakın",
                "Polimer kayıt dışı olsa da stiren vb. kayıtlı olabilir",
                "Müşteriye güvenli kullanım bilgisini verin",
            ],
        ),
    ]
    for i, (title, color, items) in enumerate(blocks):
        x = 0.4 + i * 4.25
        card(s, x, 1.5, 4.1, 5.35, WHITE)
        rect(s, x, 1.5, 4.1, 0.55, color, rounded=False)
        tb(s, x, 1.58, 4.1, 0.4, title, 16, True, WHITE, PP_ALIGN.CENTER)
        for j, item in enumerate(items):
            tb(s, x + 0.2, 2.2 + j * 0.72, 3.7, 0.65, f"•  {item}", 13, False, INK)
    footer(s, 10)

    # 11 CHECKLIST
    s = blank(prs)
    header_bar(
        s,
        "Bugün",
        "30 Eylül öncesi 12 maddelik kontrol listesi",
        "Sarı Excel satırlarını doldurduğunuzda kapsam ve son tarih otomatik çıkar",
    )
    checks = [
        ("CAS envanteri GBF bölüm 3 ile eşleşiyor.", TEAL),
        ("Karışım bileşenleri % × yıllık miktar = madde tonajı.", SKY),
        ("Aynı CAS, tüm tedarikçilerde toplanmış.", GOLD),
        ("Rol: imalatçı / ithalatçı / TT / alt kullanıcı net.", PURPLE),
        ("Tek temsilci yazısı sizi ismen kapsıyor mu bakıldı.", GREEN),
        ("Muafiyet varsa Ek numarası yazılı.", CORAL),
        ("CMR 1A/1B ve H400/H410 sınıfları tarandı.", TEAL),
        ("EÇBS + KKS hesabı ve KDU hazır.", SKY),
        ("Ön-MBDF / ortak kayıt veya bireysel gerekçe.", GOLD),
        ("Geçici kayıt şablonu + harç + ödeme teyidi.", PURPLE),
        ("Türkçe 16 bölümlü GBF KDU kimliğiyle duruyor.", GREEN),
        ("2026 / 2028 / 2030 tam kayıt bütçesi planlandı.", CORAL),
    ]
    for i, (txt, col) in enumerate(checks):
        coln, row = i % 2, i // 2
        x = 0.4 + coln * 6.45
        y = 1.48 + row * 0.85
        card(s, x, y, 6.25, 0.75, WHITE)
        oval(s, x + 0.15, y + 0.18, 0.38, 0.38, col)
        tb(s, x + 0.15, y + 0.24, 0.38, 0.28, "✓", 14, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.65, y + 0.12, 5.4, 0.52, txt, 13, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 11)

    # 12 KAYNAKLAR
    s = blank(prs)
    header_bar(
        s,
        "Kaynak",
        "Nereye bakacaksınız, kime soracaksınız?",
        "Bu rehber genel bilgilendirmedir; bağlayıcı karar KKS ve Bakanlık yazışmasıdır",
    )
    links = [
        ("Mevzuat", "KKDİK Yönetmeliği RG 23.06.2017/30105 Mükerrer\nDeğişiklik RG 23.12.2023/32408 (kayıt son tarihleri)\nUsul ve Esaslar, 05.08.2025 / 13216050 sayılı Olur"),
        ("Sistem", "Kimyasal Kayıt Sistemi (KKS) — EÇBS uygulaması\nkimyasallar.csb.gov.tr  •  Geçici kayıt şablonu ve kılavuzlar\nGBF paket programı (Usul ve Esaslar md. 16)"),
        ("İletişim", "Çevre Yönetimi Genel Müdürlüğü\nKimyasallar Yönetimi Dairesi Başkanlığı\nkimyasallar@csb.gov.tr  •  +90 312 410 10 00"),
        ("Firma içi", "KDU (Kimyasal Değerlendirme Uzmanı) ile dosya girişi\nGümrük / satın alma / üretim tonaj kanıtı\nYanındaki Excel: madde satırı = kapsam + son tarih + yapılacak"),
    ]
    for i, (title, body) in enumerate(links):
        col, row = i % 2, i // 2
        x = 0.4 + col * 6.45
        y = 1.5 + row * 2.35
        card(s, x, y, 6.25, 2.2, WHITE)
        pill(s, x + 0.2, y + 0.2, 2.15, 0.34, TEAL)
        tb(s, x + 0.2, y + 0.24, 2.15, 0.28, title, 13, True, WHITE, PP_ALIGN.CENTER)
        tb(s, x + 0.2, y + 0.7, 5.85, 1.3, body, 14, False, INK)
    tb(
        s,
        0.4,
        6.35,
        12.5,
        0.6,
        "Uyarı: İdari yaptırım 2872 sayılı Çevre Kanunu’ndandır. Bu sunum hukuki mütalaa yerine geçmez. "
        "Madde özelinde KDU ve gerektiğinde Bakanlık yardım masası teyidi alın.",
        13,
        False,
        MUTED,
    )
    footer(s, 12)

    prs.save(OUT)
    print(f"Saved {OUT} ({OUT.stat().st_size / 1024:.0f} KB)")
    return OUT


if __name__ == "__main__":
    build()
