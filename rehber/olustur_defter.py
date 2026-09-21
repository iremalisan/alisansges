#!/usr/bin/env python3
"""7 günlük doldurulabilir manifest defteri (Word)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Yedi_Gunluk_Manifest_Defteri.docx"

PLUM = RGBColor(42, 22, 58)
ROSE = RGBColor(196, 84, 118)
GOLD = RGBColor(201, 154, 74)
MUTED = RGBColor(110, 96, 118)


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_run(run, size=12, bold=False, color=PLUM, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def p(doc, text, size=12, bold=False, color=PLUM, space_after=8, align="left", italic=False):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run(run, size, bold, color, italic)
    return para


def lines(doc, n=4):
    for _ in range(n):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(10)
        run = para.add_run("________________________________________________________________")
        set_run(run, 12, False, RGBColor(200, 188, 196))


def heading_bar(doc, kicker, title):
    p(doc, kicker.upper(), 11, True, ROSE, 2)
    p(doc, title, 22, True, PLUM, 10)


THEMES = [
    ("1. gün · Netleş", "Bugün tek isteği kilitle. İlk somut adımı bugün at."),
    ("2. gün · Ritim", "Aynı cümle. En küçük görünür hareket yeter."),
    ("3. gün · Hikâye", "8 dakikalık scripting yaz. Bir kapı çal."),
    ("4. gün · Engel", "İç freni adlandır. WOOP planını sıkılaştır."),
    ("5. gün · Kimlik", "Günü bunu zaten yaşayan kişi olarak geçir."),
    ("6. gün · Kanıt", "3 küçük işaret + 1 büyükçe adım."),
    ("7. gün · Mühürle", "Cümleyi %10 netleştir. 21 güne kilit at."),
]


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    p(doc, "7 GÜNLÜK HIZLI MANIFEST DEFTERİ", 11, True, GOLD, 4, "center")
    p(doc, "Yaz, hisset, bir adım at.", 28, True, PLUM, 8, "center")
    p(
        doc,
        "Bu defter izlenecek 12 slaytlık rehberin uygulama yüzüdür. "
        "Her gün aynı iskelet: sabah 3, öğlen 6, gece 9 yazış + SATS + 1 somut adım. "
        "Yeni teknik ekleme. Mükemmel cümle arama.",
        12,
        False,
        MUTED,
        14,
        "center",
    )

    table = doc.add_table(rows=3, cols=1)
    notes = [
        "Kural: 7 gün boyunca tek istek.",
        "Kural: Cümle şimdiki zamanda, 25 kelimeyi geçmesin.",
        "Kural: Kanıt avı yok. Defteri kapat, güne dön.",
    ]
    for i, note in enumerate(notes):
        cell = table.rows[i].cells[0]
        shade(cell, "F8E8EE" if i != 1 else "EDE4F5")
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(note)
        set_run(run, 13, True, PLUM)

    p(doc, "", 8, space_after=6)
    heading_bar(doc, "Sayfa 1", "İstek kilidi")
    p(doc, "Bu 7 günde tek isteğim:", 12, True)
    lines(doc, 2)
    p(doc, "369 cümlem (şimdiki zaman, minnettarlık + duygu):", 12, True)
    lines(doc, 3)
    p(doc, "Bunu 90 günde mümkün kılan en küçük gerçek işaret:", 12, True)
    lines(doc, 2)

    doc.add_page_break()
    heading_bar(doc, "Sayfa 2", "SATS sahnesi")
    p(doc, "Sahne, dileğin gerçekleşmesinden SONRA olmalıdır. 5–10 saniye, birinci tekil şahıs, GIF gibi döngü.", 12, False, MUTED, 10)
    prompts = [
        "Neredeyim, saat kaç?",
        "Kiminle / hangi cümleyi duyuyorum?",
        "Elimde, yüzümde, odada ne var? (3 duyu)",
        "Vücutta hangi ‘tabii ki böyle’ hissi var?",
    ]
    for prompt in prompts:
        p(doc, prompt, 12, True)
        lines(doc, 2)

    doc.add_page_break()
    heading_bar(doc, "Sayfa 3", "WOOP")
    labels = [
        ("W · Wish / İstek", "Tek dilek, 369 cümlesiyle aynı."),
        ("O · Outcome / Sonuç", "Olduğunda en güzel 30 saniye."),
        ("O · Obstacle / Engel", "Dış dünya değil: içindeki asıl fren."),
        ("P · Plan", "Eğer [engel], o zaman [2 dakikalık davranış]."),
    ]
    for title, hint in labels:
        p(doc, title, 14, True, ROSE, 2)
        p(doc, hint, 11, False, MUTED, 4, italic=True)
        lines(doc, 3)

    for day, theme in THEMES:
        doc.add_page_break()
        heading_bar(doc, day.split("·")[0].strip(), day)
        p(doc, theme, 12, False, MUTED, 10, italic=True)

        grid = doc.add_table(rows=4, cols=2)
        grid.style = "Table Grid"
        cells = [
            ("Sabah ×3 bitti mi?", "Evet / Hayır    saat: ______"),
            ("Öğlen ×6 bitti mi?", "Evet / Hayır    saat: ______"),
            ("Gece ×9 bitti mi?", "Evet / Hayır    saat: ______"),
            ("SATS ile mi uyudum?", "Evet / Hayır    his (1–10): ______"),
        ]
        for i, (a, b) in enumerate(cells):
            grid.rows[i].cells[0].text = a
            grid.rows[i].cells[1].text = b
            shade(grid.rows[i].cells[0], "EDE4F5")
            for cell in grid.rows[i].cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        set_run(run, 12, True, PLUM)

        p(doc, "", 6, space_after=8)
        p(doc, "Bugünün tek somut adımı (10 dakikayı geçmesin):", 12, True)
        lines(doc, 2)
        p(doc, "Adımı attım mı, ne oldu?", 12, True)
        lines(doc, 2)
        p(doc, "Şüphe geldiyse yaz ve bırak (tartışma):", 12, True)
        lines(doc, 2)
        p(doc, "Bugün ‘oldu’ hissine en çok yaklaştığım an:", 12, True)
        lines(doc, 2)
        if "3. gün" in day:
            p(doc, "Scripting (olmuş bir salı sabahı, 8 dakika):", 12, True, ROSE)
            lines(doc, 8)
        if "7. gün" in day:
            p(doc, "7 günün özeti ve önümüzdeki 21 güne kilit cümle:", 12, True, ROSE)
            lines(doc, 6)

    p(doc, "", 8)
    p(
        doc,
        "Kapanış: Bir istek. Aynı cümle. Aynı sahne. Her gün bir adım. 7 günden sonra yeni teknik arama.",
        12,
        True,
        PLUM,
        6,
        "center",
    )
    p(doc, "Kaynak omurga: Neville Goddard — Feeling is the Secret  •  woopmylife.org  •  369 yazma ritmi", 10, False, MUTED, 0, "center")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
