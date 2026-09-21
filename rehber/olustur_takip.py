#!/usr/bin/env python3
"""7 günlük manifest takip Excel’i."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Yedi_Gunluk_Manifest_Takip.xlsx"

PLUM = "2A163A"
ROSE = "C45476"
GOLD = "C99A4A"
CREAM = "FBF6EE"
LILAC = "EDE4F5"
YELLOW = "FFF3B0"
WHITE = "FFFFFF"
GREEN = "2E8060"
MUTED = "6E6076"
INK = "20182A"

thin = Border(
    left=Side(style="thin", color="D8CCE0"),
    right=Side(style="thin", color="D8CCE0"),
    top=Side(style="thin", color="D8CCE0"),
    bottom=Side(style="thin", color="D8CCE0"),
)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(bold=False, color=INK, size=11, italic=False) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)


def align(h="left", v="center", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def style_range(ws, cells, **kwargs):
    for row in ws[cells] if ":" in cells else [ws[cells]]:
        iterable = row if isinstance(row, tuple) else (row,)
        for cell in iterable:
            for k, v in kwargs.items():
                setattr(cell, k, v)


def col_widths(ws, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def header_row(ws, row, values, fill_color=PLUM, font_color=WHITE):
    for i, value in enumerate(values, 1):
        cell = ws.cell(row, i, value)
        cell.fill = fill(fill_color)
        cell.font = font(True, font_color, 11)
        cell.alignment = align("center")
        cell.border = thin


def build() -> Path:
    wb = Workbook()

    # --- TALIMAT ---
    ws = wb.active
    ws.title = "Nasıl kullanılır"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ROSE
    ws.merge_cells("B2:G2")
    ws["B2"] = "7 günde hızlı manifest — takip tablosu"
    ws["B2"].font = font(True, PLUM, 22)
    ws.merge_cells("B3:G3")
    ws["B3"] = "Sarı hücrelere yaz. Diğer sütunlar otomatik dolar. 7 gün aynı istek, aynı cümle."
    ws["B3"].font = font(False, MUTED, 12)

    steps = [
        ("1", "Cümle sayfasına 369 cümleni yaz. Süzgeç 5/5 olana kadar küçült veya somutlaştır."),
        ("2", "WOOP sayfasında iç engeli ve eğer-o zaman planını doldur."),
        ("3", "Her gün Takip sayfasında Evet/Hayır işaretle, somut adımı bir cümleyle yaz, his notunu 1–10 ver."),
        ("4", "Skor %80 altındaysa teknik ekleme; eksik katmanı (yazış / SATS / adım) tamamla."),
        ("5", "7. gün özet satırını oku. Cümleyi gerekirse %10 netleştir, 21 güne uzat."),
    ]
    header_row(ws, 5, ["", "Adım", "Ne yapacaksın", "", "", "", ""])
    ws.merge_cells("C5:G5")
    for i, (num, text) in enumerate(steps):
        r = 6 + i
        ws.cell(r, 2, num).fill = fill(ROSE)
        ws.cell(r, 2).font = font(True, WHITE, 14)
        ws.cell(r, 2).alignment = align("center")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ws.cell(r, 3, text).font = font(False, INK, 12)
        ws.cell(r, 3).alignment = align("left")
        ws.cell(r, 3).fill = fill(CREAM if i % 2 == 0 else LILAC)
        ws.row_dimensions[r].height = 36
        for c in range(2, 8):
            ws.cell(r, c).border = thin

    ws.merge_cells("B12:G12")
    ws["B12"] = "Hız formülü: tek istek + 369 (3-6-9) + gece SATS + her gün 1 somut adım (WOOP)."
    ws["B12"].font = font(True, PLUM, 13)
    ws["B12"].fill = fill(YELLOW)
    ws["B12"].alignment = align("left")
    ws.row_dimensions[12].height = 28

    ws.merge_cells("B14:G16")
    ws["B14"] = (
        "Dürüst not: Bu tablo bir sihir garantisi değildir. Dikkatini, hissini ve 10 dakikalık davranışını aynı yöne kilitler. "
        "Kaynak omurga: Neville Goddard — Feeling is the Secret (SATS); Gabriele Oettingen — WOOP (woopmylife.org); "
        "elle yazılan 369 tekrarı bir odak ritmi olarak kullanılır, Tesla formülü değildir."
    )
    ws["B14"].font = font(False, MUTED, 11)
    ws["B14"].alignment = align("left")
    col_widths(ws, {"A": 3, "B": 8, "C": 22, "D": 18, "E": 18, "F": 18, "G": 28})
    ws.row_dimensions[2].height = 28
    ws.freeze_panes = "B6"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    # --- CUMLE ---
    c = wb.create_sheet("Cümle")
    c.sheet_properties.tabColor = GOLD
    c.sheet_view.showGridLines = False
    c.merge_cells("B2:F2")
    c["B2"] = "369 cümlesi"
    c["B2"].font = font(True, PLUM, 20)
    c.merge_cells("B3:F3")
    c["B3"] = "Sarı hücre senin. Şimdiki zaman, minnettarlık, duygu, 25 kelimeyi geçme."
    c["B3"].font = font(False, MUTED, 12)

    labels = [
        (5, "Tek isteğim (ham hali)"),
        (7, "369 cümlem (kilit)"),
        (9, "SATS sahnesi (5–10 sn, olduktan SONRA)"),
        (11, "Bugün atabileceğim en küçük adım"),
    ]
    for row, label in labels:
        c.cell(row, 2, label).font = font(True, ROSE, 12)
        c.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=6)
        cell = c.cell(row + 1, 2, "")
        cell.fill = fill(YELLOW)
        cell.alignment = align("left")
        cell.border = thin
        c.row_dimensions[row + 1].height = 48
        for col in range(2, 7):
            c.cell(row + 1, col).fill = fill(YELLOW)
            c.cell(row + 1, col).border = thin

    c["B6"] = "ör. bu ay nakit akışım rahatlasın"
    c["B6"].font = font(False, MUTED, 11, True)
    c["B8"] = "ör. Çok minnettarım, bu ay faturalarım rahat ödeniyor ve kendimi güvende hissediyorum."
    c["B8"].font = font(False, MUTED, 11, True)
    c["B10"] = "ör. Mutfakta kahvemi içerken bakiyeyi görüp omuzlarımın düştüğünü hissediyorum."
    c["B10"].font = font(False, MUTED, 11, True)
    c["B12"] = "ör. Bugün 1 teklif / 1 mail / 1 fiyat / 1 yürüyüş"
    c["B12"].font = font(False, MUTED, 11, True)

    header_row(c, 14, ["", "Süzgeç", "Evet / Hayır", "Not", "", ""])
    checks = [
        "90 gün içinde mümkün mü?",
        "Hayal edince vücutta yumuşama var mı?",
        "Atabileceğin en az bir adım var mı?",
        "Başkasının iradesini zorlamıyor mu?",
        "Ölçülebilir mi? (tarih, tutar, olay)",
    ]
    dv_yn = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
    dv_yn.error = "Evet veya Hayır seç"
    dv_yn.errorTitle = "Geçersiz"
    dv_yn.prompt = "Evet / Hayır"
    dv_yn.promptTitle = "Süzgeç"
    c.add_data_validation(dv_yn)
    for i, q in enumerate(checks):
        r = 15 + i
        c.cell(r, 2, q).font = font(False, INK, 12)
        c.cell(r, 2).fill = fill(LILAC)
        c.cell(r, 3, "").fill = fill(YELLOW)
        c.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        c.cell(r, 4).fill = fill(YELLOW)
        for col in range(2, 7):
            c.cell(r, col).border = thin
            c.cell(r, col).alignment = align("left")
        dv_yn.add(c.cell(r, 3))
        c.row_dimensions[r].height = 28
    c.merge_cells("B21:F21")
    c["B21"] = '=IF(COUNTIF(C15:C19,"Evet")=5,"Süzgeç tamam — 7 güne kilit at.","Süzgeç eksik — cümleyi küçült veya somutlaştır.")'
    c["B21"].font = font(True, PLUM, 13)
    c["B21"].fill = fill(CREAM)
    col_widths(c, {"A": 3, "B": 52, "C": 16, "D": 22, "E": 18, "F": 18})
    c.freeze_panes = "B5"

    # --- WOOP ---
    w = wb.create_sheet("WOOP")
    w.sheet_properties.tabColor = "7A5CA8"
    w.sheet_view.showGridLines = False
    w.merge_cells("B2:E2")
    w["B2"] = "WOOP — istek, sonuç, iç engel, eğer-o zaman"
    w["B2"].font = font(True, PLUM, 20)
    w.merge_cells("B3:E3")
    w["B3"] = "Gabriele Oettingen. Tek başına olumlu hayal eylemi azaltabilir. Engel + plan hız kazandırır."
    w["B3"].font = font(False, MUTED, 12)

    woop_rows = [
        (5, "W", "Wish / İstek", "Cümle!B8", "Tek dilek, 369 cümlesiyle aynı."),
        (10, "O", "Outcome / Sonuç", "", "Olduğunda en güzel 30 saniye. Göğüste ne açılıyor?"),
        (15, "O", "Obstacle / İç engel", "", "Erteleme, utanç, hak etmiyorum, gece telefon… Dış dünya değil."),
        (20, "P", "Plan / Eğer-o zaman", "", "Eğer [engel belirirse], o zaman [2 dakikalık davranış]."),
    ]
    for row, letter, title, formula, hint in woop_rows:
        w.cell(row, 2, letter).fill = fill(ROSE if letter != "P" else GOLD)
        w.cell(row, 2).font = font(True, WHITE, 18)
        w.cell(row, 2).alignment = align("center")
        w.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        w.cell(row, 3, title).font = font(True, PLUM, 14)
        w.cell(row, 3).fill = fill(LILAC)
        w.merge_cells(start_row=row + 1, start_column=2, end_row=row + 2, end_column=5)
        cell = w.cell(row + 1, 2, "")
        if formula:
            cell.value = f"={formula}"
        cell.fill = fill(YELLOW)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for r in range(row, row + 3):
            for col in range(2, 6):
                w.cell(r, col).border = thin
                if r > row:
                    w.cell(r, col).fill = fill(YELLOW)
        w.cell(row + 3, 2, hint).font = font(False, MUTED, 10, True)
        w.merge_cells(start_row=row + 3, start_column=2, end_row=row + 3, end_column=5)
        w.row_dimensions[row + 1].height = 22
        w.row_dimensions[row + 2].height = 28
    w.merge_cells("B25:E26")
    w["B25"] = (
        "Hazır iskelet: Eğer öğleden sonra “nasıl olsa olmaz” diye kaydırırsam, "
        "o zaman telefonu başka odaya koyar 2 dakikalık adımı hemen bitiririm."
    )
    w["B25"].font = font(False, INK, 12)
    w["B25"].fill = fill("E8F4EC")
    w["B25"].alignment = align("left")
    col_widths(w, {"A": 3, "B": 10, "C": 28, "D": 28, "E": 36})

    # --- TAKIP ---
    t = wb.create_sheet("Takip")
    t.sheet_properties.tabColor = GREEN
    t.freeze_panes = "C5"
    t.merge_cells("B2:L2")
    t["B2"] = "7 günlük ritim — sarı hücreleri doldur"
    t["B2"].font = font(True, PLUM, 20)
    t.merge_cells("B3:L3")
    t["B3"] = '= "Kilit cümle: " & IF(Cümle!B8="","(Cümle sayfasına yaz)",Cümle!B8)'
    t["B3"].font = font(True, ROSE, 12)
    t["B3"].alignment = align("left")

    headers = [
        "Gün",
        "Tema",
        "Tarih",
        "Sabah ×3",
        "Öğlen ×6",
        "Gece ×9",
        "SATS",
        "Somut adım (ne yaptım?)",
        "Adım oldu mu?",
        "His 1–10",
        "Gün skoru",
        "Not / şüpheyi bırak",
    ]
    header_row(t, 5, headers)
    themes = ["Netleş", "Ritim", "Hikâye", "Engel", "Kimlik", "Kanıt", "Mühürle"]
    dv_yn2 = DataValidation(type="list", formula1='"Evet,Hayır"', allow_blank=True)
    dv_his = DataValidation(type="whole", operator="between", formula1="1", formula2="10", allow_blank=True)
    dv_his.error = "1 ile 10 arası yaz"
    t.add_data_validation(dv_yn2)
    t.add_data_validation(dv_his)
    for i, theme in enumerate(themes):
        r = 6 + i
        t.cell(r, 1, i + 1).font = font(True, WHITE, 14)
        t.cell(r, 1).fill = fill(ROSE)
        t.cell(r, 1).alignment = align("center")
        t.cell(r, 2, theme).font = font(True, PLUM, 12)
        t.cell(r, 2).fill = fill(LILAC)
        for col in (3, 4, 5, 6, 7, 8, 9, 10, 12):
            t.cell(r, col).fill = fill(YELLOW)
            t.cell(r, col).border = thin
            t.cell(r, col).alignment = align("center" if col != 8 and col != 12 else "left")
        t.cell(r, 11, f'=IF(COUNTA(D{r}:G{r},I{r})=0,"—",INT(20*(COUNTIF(D{r},"Evet")+COUNTIF(E{r},"Evet")+COUNTIF(F{r},"Evet")+COUNTIF(G{r},"Evet")+COUNTIF(I{r},"Evet"))))')
        t.cell(r, 11).font = font(True, PLUM, 14)
        t.cell(r, 11).alignment = align("center")
        t.cell(r, 11).number_format = '0"%"'
        for col in range(1, 13):
            t.cell(r, col).border = thin
        t.row_dimensions[r].height = 42
        dv_yn2.add(t.cell(r, 4))
        dv_yn2.add(t.cell(r, 5))
        dv_yn2.add(t.cell(r, 6))
        dv_yn2.add(t.cell(r, 7))
        dv_yn2.add(t.cell(r, 9))
        dv_his.add(t.cell(r, 10))

    t.conditional_formatting.add(
        "K6:K12",
        ColorScaleRule(start_type="num", start_value=0, start_color="F8E8EE", mid_type="num", mid_value=60, mid_color="FFF3B0", end_type="num", end_value=100, end_color="8FCB9B"),
    )

    t.merge_cells("B14:C14")
    t["B14"] = "7 gün ortalama skor"
    t["B14"].font = font(True, WHITE, 12)
    t["B14"].fill = fill(PLUM)
    t.merge_cells("D14:E14")
    t["D14"] = '=IF(COUNT(K6:K12)=0,"—",INT(AVERAGE(K6:K12)))'
    t["D14"].font = font(True, PLUM, 16)
    t["D14"].fill = fill(GOLD)
    t["D14"].alignment = align("center")
    t["D14"].number_format = '0"%"'
    t.merge_cells("F14:L14")
    t["F14"] = '=IF(D14="—","Henüz gün işaretlenmedi.",IF(D14>=80,"Ritim oturmuş. 21 güne uzat, yeni teknik ekleme.","Eksik katmanı tamamla: yazış, SATS veya somut adım."))'
    t["F14"].font = font(True, PLUM, 12)
    t["F14"].fill = fill(CREAM)

    t.merge_cells("B16:L16")
    t["B16"] = "Evet sayısı (yazış + SATS + adım)"
    t["B16"].font = font(True, MUTED, 11)
    for i, label in enumerate(["Sabah", "Öğlen", "Gece", "SATS", "Adım"]):
        t.cell(17, 2 + i, label).font = font(True, WHITE, 11)
        t.cell(17, 2 + i).fill = fill(PLUM)
        t.cell(17, 2 + i).alignment = align("center")
    t["B18"] = '=COUNTIF(D6:D12,"Evet")'
    t["C18"] = '=COUNTIF(E6:E12,"Evet")'
    t["D18"] = '=COUNTIF(F6:F12,"Evet")'
    t["E18"] = '=COUNTIF(G6:G12,"Evet")'
    t["F18"] = '=COUNTIF(I6:I12,"Evet")'
    for col in range(2, 7):
        t.cell(18, col).font = font(True, PLUM, 14)
        t.cell(18, col).alignment = align("center")
        t.cell(18, col).fill = fill(LILAC)
        t.cell(18, col).border = thin

    chart = BarChart()
    chart.type = "col"
    chart.title = "Katman başına Evet (7 gün)"
    chart.y_axis.title = "gün"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 7
    data = Reference(t, min_col=2, min_row=17, max_col=6, max_row=18)
    cats = Reference(t, min_col=2, min_row=17, max_col=6)
    chart.add_data(data, from_rows=True, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.style = 10
    chart.width = 15
    chart.height = 7
    t.add_chart(chart, "B20")

    col_widths(
        t,
        {
            "A": 8,
            "B": 14,
            "C": 14,
            "D": 13,
            "E": 13,
            "F": 13,
            "G": 12,
            "H": 36,
            "I": 14,
            "J": 12,
            "K": 12,
            "L": 32,
        },
    )
    t.auto_filter.ref = "A5:L12"
    t.page_setup.orientation = "landscape"
    t.page_setup.fitToPage = True
    t.page_setup.fitToWidth = 1
    t.page_setup.fitToHeight = 1
    t.print_title_rows = "1:5"
    t.page_setup.paperSize = t.PAPERSIZE_A4

    # --- KAYNAKLAR ---
    k = wb.create_sheet("Kaynaklar")
    k.sheet_properties.tabColor = "C99A4A"
    k.sheet_view.showGridLines = False
    k.merge_cells("B2:E2")
    k["B2"] = "Nereleri araştır, nerede takılma"
    k["B2"].font = font(True, PLUM, 20)
    header_row(k, 4, ["", "Öncelik", "Kaynak", "Ne için", "Not"])
    rows = [
        ("Bu hafta", "Neville Goddard — Feeling is the Secret", "SATS, uykunun eşiği, his", "Kısa. Önce bunu bitir."),
        ("Bu hafta", "Neville Goddard — The Power of Awareness", "Kimlik / varsayım / sonda yaşamak", "Cümleyi kimlik haline getirir."),
        ("Bu hafta", "woopmylife.org", "Wish-Outcome-Obstacle-Plan", "5 dakikalık resmi araç."),
        ("Bu hafta", "Gollwitzer — implementation intentions", "Eğer-o zaman planlarının etkisi", "d ≈ 0.65; niyeti eyleme bağlar."),
        ("Sonra", "Gabriele Oettingen — Rethinking Positive Thinking", "Neden sadece hayal yetmez", "WOOP’un kitabı."),
        ("Sonra", "Neville — The Law and the Promise", "Uygulama hikâyeleri", "Teknik oturduktan sonra."),
        ("Sonra", "James Clear — Atomic Habits", "Küçük günlük adım", "WOOP’un davranış yüzü."),
        ("Uzak dur", "Sonsuz koç reeli / anında zengin ol vaadi", "Dağıtır", "7 gün protokolü bozar."),
        ("Uzak dur", "Başkasının iradesini manipüle etme teknikleri", "Zararlı ve boşa", "Karşılıklı, serbest hali iste."),
        ("Uzak dur", "Aynı anda 10 yöntem", "Odak kaçması", "Tek protokol, 7 gün."),
    ]
    for i, (prio, src, why, note) in enumerate(rows):
        r = 5 + i
        k.cell(r, 2, prio).font = font(True, WHITE if prio != "Uzak dur" else WHITE, 11)
        k.cell(r, 2).fill = fill(GREEN if prio == "Bu hafta" else (GOLD if prio == "Sonra" else ROSE))
        k.cell(r, 2).alignment = align("center")
        k.cell(r, 3, src).font = font(False, INK, 11)
        k.cell(r, 4, why).font = font(False, INK, 11)
        k.cell(r, 5, note).font = font(False, MUTED, 11)
        bg = CREAM if i % 2 == 0 else LILAC
        for col in range(3, 6):
            k.cell(r, col).fill = fill(bg)
        for col in range(2, 6):
            k.cell(r, col).border = thin
            k.cell(r, col).alignment = align("left")
        k.row_dimensions[r].height = 32
    col_widths(k, {"A": 3, "B": 14, "C": 48, "D": 42, "E": 36})
    k.freeze_panes = "B5"

    # sample data so the file is not an empty skeleton
    c["B6"] = "Bu ay nakit akışım rahatlasın"
    c["B8"] = "Çok minnettarım, bu ay faturalarım rahat ödeniyor ve kendimi güvende hissediyorum."
    c["B10"] = "Mutfakta kahvemi içerken bakiyeyi görüp omuzlarımın düştüğünü hissediyorum."
    c["B12"] = "Bugün bir teklif maili gönderiyorum."
    for r in range(15, 20):
        c.cell(r, 3, "Evet")
    w["B11"] = "Faturalar ödenmiş, omuzlar düşük, nefes rahat; ‘artık böyle’ hissi."
    w["B16"] = "Öğleden sonra ‘nasıl olsa olmaz’ deyip telefonu kaydırmak."
    w["B21"] = "Eğer kaydırmaya başlarsam, telefonu başka odaya koyar 2 dakikalık maili hemen bitiririm."
    t["C6"] = "21.09.2026"
    t["D6"] = "Evet"
    t["E6"] = "Evet"
    t["F6"] = "Evet"
    t["G6"] = "Evet"
    t["H6"] = "Teklif maili taslağını bitirdim."
    t["I6"] = "Evet"
    t["J6"] = 8
    t["L6"] = "Şüphe geldi, yazıp kapattım."

    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
