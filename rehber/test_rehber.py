#!/usr/bin/env python3
"""Hızlı manifest paketinin temel bütünlük testleri."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from olustur_defter import build as build_defter
from olustur_gorsel import SLIDES, build as build_gorsel
from olustur_sunum import TOTAL, build as build_sunum
from olustur_takip import build as build_takip

ROOT = Path(__file__).resolve().parent


def test_sunum_has_twelve_slides_and_key_phrases():
    path = build_sunum()
    assert path.exists() and path.stat().st_size > 40_000
    prs = Presentation(path)
    assert len(prs.slides) == TOTAL == 12
    blob = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    for needle in (
        "7 günde hızlı",
        "369",
        "SATS",
        "WOOP",
        "Feeling is the Secret",
        "woopmylife.org",
        "tek istek",
        "20 dakikada başla",
    ):
        assert needle in blob, needle


def test_gorsel_pdf_and_pngs():
    pdf, pngs = build_gorsel()
    assert pdf.exists() and pdf.stat().st_size > 80_000
    assert len(pngs) == 12 == len(SLIDES)
    for path in pngs:
        assert path.exists() and path.stat().st_size > 8_000
        assert path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


def test_defter_is_docx_with_seven_days():
    path = build_defter()
    assert path.exists() and path.stat().st_size > 20_000
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    for needle in ("369 cümlem", "SATS", "WOOP", "1. gün", "7. gün", "Scripting"):
        assert needle in xml, needle


def test_takip_workbook_formulas_and_sheets():
    path = build_takip()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Nasıl kullanılır", "Cümle", "WOOP", "Takip", "Kaynaklar"]
    takip = wb["Takip"]
    assert takip["K6"].value.startswith("=IF(COUNTA")
    assert takip["D14"].value.startswith("=IF(COUNT")
    assert takip["B18"].value == '=COUNTIF(D6:D12,"Evet")'
    assert wb["Cümle"]["B21"].value.startswith("=IF(COUNTIF")
    assert "Feeling is the Secret" in wb["Kaynaklar"]["C5"].value
    assert wb["Takip"]["D6"].value == "Evet"
    assert len(takip._charts) == 1


if __name__ == "__main__":
    test_sunum_has_twelve_slides_and_key_phrases()
    test_gorsel_pdf_and_pngs()
    test_defter_is_docx_with_seven_days()
    test_takip_workbook_formulas_and_sheets()
    print("ok")
