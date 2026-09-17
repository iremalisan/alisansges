from __future__ import annotations

import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_kkdik_kapsam_tablosu import (
    EXAMPLES,
    EXEMPT,
    FIRST_DATA_ROW,
    IN_SCOPE,
    NO_REG_ART,
    NO_REG_DU,
    NO_REG_OR,
    NO_REG_TON,
    NOTIFY,
    create_workbook,
    evaluate_row,
    full_deadline,
    substance_tonnage,
)


class LogicTests(unittest.TestCase):
    def test_mixture_tonnage_uses_percentage(self) -> None:
        self.assertAlmostEqual(substance_tonnage("Karışım", 12, "ton", 40), 4.8)
        self.assertAlmostEqual(substance_tonnage("Madde", 400, "kg", 100), 0.4)
        self.assertAlmostEqual(substance_tonnage("Madde", 20, "ton", 100), 20)

    def test_importer_over_one_ton_is_in_scope(self) -> None:
        result = evaluate_row("Karışım", "İthalatçı", 12, "ton", 40, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], IN_SCOPE)
        self.assertEqual(result["deadline"], "31.12.2030")
        self.assertEqual(result["provisional"], "30.09.2026")

    def test_downstream_user_has_no_registration(self) -> None:
        result = evaluate_row("Madde", "Alt kullanıcı", 20, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], NO_REG_DU)
        self.assertEqual(result["deadline"], "")

    def test_below_one_ton_is_out(self) -> None:
        result = evaluate_row("Madde", "İthalatçı", 400, "kg", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], NO_REG_TON)

    def test_only_representative_covers_importer(self) -> None:
        result = evaluate_row("Madde", "İthalatçı", 50, "ton", 100, "Evet", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], NO_REG_OR)

    def test_cmr_pulls_deadline_to_2026(self) -> None:
        result = evaluate_row("Madde", "İmalatçı", 2, "ton", 100, "Hayır", "Yok", "Evet", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], IN_SCOPE)
        self.assertEqual(result["deadline"], "31.12.2026")

    def test_aquatic_100t_is_2026(self) -> None:
        result = evaluate_row("Madde", "İthalatçı", 120, "ton", 100, "Hayır", "Yok", "Hayır", "Evet", "Hayır", "Hayır")
        self.assertEqual(result["deadline"], "31.12.2026")

    def test_100_to_1000_without_special_class_is_2028(self) -> None:
        result = evaluate_row("Madde", "İthalatçı", 150, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], IN_SCOPE)
        self.assertEqual(result["deadline"], "31.12.2028")

    def test_1000_ton_is_2026(self) -> None:
        self.assertEqual(full_deadline(1000, "Hayır", "Hayır"), "31.12.2026")
        result = evaluate_row("Madde", "İmalatçı", 1500, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["deadline"], "31.12.2026")

    def test_article_without_intended_release(self) -> None:
        result = evaluate_row(
            "Eşya", "Eşya üreticisi / ithalatçısı", 800, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Hayır"
        )
        self.assertEqual(result["status"], NO_REG_ART)

    def test_article_svhc_notification(self) -> None:
        result = evaluate_row(
            "Eşya", "Eşya üreticisi / ithalatçısı", 5, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Hayır", "Evet"
        )
        self.assertEqual(result["status"], NOTIFY)

    def test_article_with_intended_release_is_in_scope(self) -> None:
        result = evaluate_row(
            "Eşya", "Eşya üreticisi / ithalatçısı", 3, "ton", 100, "Hayır", "Yok", "Hayır", "Hayır", "Evet", "Hayır"
        )
        self.assertEqual(result["status"], IN_SCOPE)
        self.assertEqual(result["provisional"], "30.09.2026")

    def test_exemption_wins(self) -> None:
        result = evaluate_row("Madde", "İthalatçı", 30, "ton", 100, "Hayır", "Ek-5 muafiyet", "Hayır", "Hayır", "Hayır", "Hayır")
        self.assertEqual(result["status"], EXEMPT)
        self.assertEqual(result["provisional"], "")


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.path = Path("/tmp/kkdik-kapsam-test.xlsx")
        create_workbook(cls.path)
        cls.wb = load_workbook(cls.path)

    def test_expected_sheets_exist(self) -> None:
        self.assertEqual(
            self.wb.sheetnames,
            [
                "Madde Envanteri",
                "Rolüne Göre Ne Yapmalı",
                "Sistem ve Takvim",
                "Muafiyetler",
                "Nasıl Kullanılır",
            ],
        )

    def test_example_rows_match_logic(self) -> None:
        ws = self.wb["Madde Envanteri"]
        self.assertEqual(len(EXAMPLES), 8)
        self.assertEqual(ws["A11"].value, EXAMPLES[0][0])
        self.assertEqual(ws["D11"].value, "İthalatçı")
        self.assertIn(IN_SCOPE, str(ws["Q11"].value))
        self.assertIn("30.09.2026", str(ws["S11"].value))
        self.assertEqual(ws["A12"].value, EXAMPLES[1][0])
        self.assertIn(NO_REG_DU, str(ws["Q12"].value))

    def test_formulas_cover_status_and_deadline(self) -> None:
        ws = self.wb["Madde Envanteri"]
        self.assertTrue(str(ws["P11"].value).startswith("="))
        self.assertTrue(str(ws["Q11"].value).startswith("="))
        self.assertTrue(str(ws["R11"].value).startswith("="))
        self.assertTrue(str(ws["T11"].value).startswith("="))
        self.assertIn("kg", str(ws["P11"].value))
        self.assertIn(IN_SCOPE, str(ws["Q11"].value))

    def test_file_is_unlocked(self) -> None:
        import zipfile

        z = zipfile.ZipFile(self.path)
        xml = z.read("xl/workbook.xml").decode("utf-8")
        self.assertNotIn("workbookProtection", xml)
        self.assertIn("fullCalcOnLoad", xml)
        for name in z.namelist():
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                sheet = z.read(name).decode("utf-8")
                self.assertNotIn("sheetProtection", sheet)
        for ws in self.wb.worksheets:
            self.assertFalse(ws.protection.sheet)

    def test_cached_example_status_present(self) -> None:
        import zipfile

        xml = zipfile.ZipFile(self.path).read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertIn(IN_SCOPE, xml)
        self.assertIn(NO_REG_DU, xml)
        self.assertIn("30.09.2026", xml)
        self.assertNotIn("<v />", xml)


if __name__ == "__main__":
    unittest.main()
