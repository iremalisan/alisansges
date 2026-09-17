from __future__ import annotations

import re
import unittest
from pathlib import Path

from openpyxl import load_workbook

from create_su_verimliligi_kapsam_tablosu import EK2_NACE, create_workbook


def normalize_nace(value: object) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) < 4:
        return "GEÇERSİZ"
    return f"{digits[:2]}.{digits[2:4]}"


def scope_status(nace: object, employees: object) -> str:
    normalized = normalize_nace(nace)
    in_list = normalized in {code for code, *_ in EK2_NACE}
    enough_staff = isinstance(employees, (int, float)) and employees >= 50
    if in_list and enough_staff:
        return "KAPSAMDA"
    return "KAPSAM DIŞI"


class Ek2ListTests(unittest.TestCase):
    def test_list_has_148_unique_codes(self) -> None:
        codes = [item[0] for item in EK2_NACE]
        self.assertEqual(len(codes), 148)
        self.assertEqual(len(set(codes)), 148)

    def test_codes_are_four_digit_nace(self) -> None:
        for code, *_ in EK2_NACE:
            self.assertRegex(code, r"^\d{2}\.\d{2}$")


class ScopeLogicTests(unittest.TestCase):
    def test_listed_nace_with_50_or_more_is_in_scope(self) -> None:
        self.assertEqual(scope_status("10.11", 50), "KAPSAMDA")
        self.assertEqual(scope_status("10.11.01", 120), "KAPSAMDA")
        self.assertEqual(scope_status("10,11", 51), "KAPSAMDA")
        self.assertEqual(scope_status("1011", 80), "KAPSAMDA")

    def test_listed_nace_under_50_is_out_of_scope(self) -> None:
        self.assertEqual(scope_status("10.11", 49), "KAPSAM DIŞI")
        self.assertEqual(scope_status("10.11", 30), "KAPSAM DIŞI")

    def test_unlisted_nace_is_out_of_scope_even_with_many_employees(self) -> None:
        self.assertEqual(scope_status("47.11", 200), "KAPSAM DIŞI")
        self.assertEqual(scope_status("85.20", 500), "KAPSAM DIŞI")

    def test_known_ek2_codes_from_regulation(self) -> None:
        codes = {item[0] for item in EK2_NACE}
        for code in ("01.41", "10.81", "20.16", "24.10", "35.11", "42.12"):
            self.assertIn(code, codes)


class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.path = Path("/tmp/su-verimliligi-kapsam-test.xlsx")
        create_workbook(cls.path)
        cls.wb = load_workbook(cls.path)

    def test_expected_sheets_exist(self) -> None:
        self.assertEqual(
            self.wb.sheetnames,
            ["Firma Kontrolü", "Ek-2 NACE Listesi", "Nasıl Kullanılır"],
        )

    def test_example_rows_and_formulas(self) -> None:
        ws = self.wb["Firma Kontrolü"]
        self.assertEqual(ws["A10"].value, "Örnek Gıda A.Ş.")
        self.assertEqual(ws["B10"].value, "10.11")
        self.assertEqual(ws["C10"].value, 120)
        self.assertIn("KAPSAMDA", ws["F10"].value)
        self.assertIn("KAPSAM DIŞI", ws["F11"].value)
        self.assertEqual(ws["C12"].value, 30)

    def test_conditional_formatting_covers_scope_column(self) -> None:
        ws = self.wb["Firma Kontrolü"]
        rules = list(ws.conditional_formatting._cf_rules.values())
        flat = [rule for group in rules for rule in group]
        formulas = [getattr(rule, "formula", None) for rule in flat]
        self.assertTrue(any(f == ['"KAPSAMDA"'] for f in formulas))
        self.assertTrue(any(f == ['"KAPSAM DIŞI"'] for f in formulas))

    def test_nace_sheet_row_count(self) -> None:
        ws = self.wb["Ek-2 NACE Listesi"]
        self.assertEqual(ws["B2"].value, "01.41")
        self.assertEqual(ws["B149"].value, "42.12")
        self.assertIsNone(ws["B150"].value)


if __name__ == "__main__":
    unittest.main()
