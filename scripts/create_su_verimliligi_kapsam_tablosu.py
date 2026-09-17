#!/usr/bin/env python3
"""Su Verimliliği Yönetmeliği Ek-2 kapsam kontrol Excel tablosunu üretir."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_PATH = Path(__file__).resolve().parents[1] / "Su_Verimliligi_Kapsam_Kontrol.xlsx"

# 27.12.2024 tarihli Su Verimliliği Yönetmeliği Ek-2 NACE kodu listesi (148 kod)
EK2_NACE: list[tuple[str, str, str, str]] = [
    ("01.41", "Sütü sağılan büyük baş hayvan yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("01.42", "Diğer sığır ve manda yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("01.43", "At ve at benzeri diğer hayvan yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("01.45", "Koyun ve keçi yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("01.47", "Kümes hayvanları yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("01.49", "Diğer hayvan yetiştiriciliği", "01", "Bitkisel ve hayvansal üretim ile avcılık ve ilgili hizmet faaliyetleri"),
    ("03.12", "Tatlı su balık yetiştiriciliği (İç su ürünleri dâhil)", "03", "Balıkçılık ve su ürünleri yetiştiriciliği"),
    ("05.10", "Taş kömürü madenciliği", "05", "Kömür ve linyit çıkartılması"),
    ("05.20", "Linyit madenciliği", "05", "Kömür ve linyit çıkartılması"),
    ("07.10", "Demir cevherleri madenciliği", "07", "Metal cevherleri madenciliği"),
    ("07.29", "Diğer demir dışı metal cevherleri madenciliği", "07", "Metal cevherleri madenciliği"),
    ("08.91", "Kimyasal ve gübreleme amaçlı mineral madenciliği", "08", "Diğer madencilik ve taş ocakçılığı"),
    ("08.93", "Tuz çıkarımı", "08", "Diğer madencilik ve taş ocakçılığı"),
    ("09.10", "Petrol ve doğal gaz çıkarımını destekleyici faaliyetler", "09", "Madenciliği destekleyici hizmet faaliyetleri"),
    ("10.11", "Etin işlenmesi ve saklanması", "10", "Gıda ürünlerinin imalatı"),
    ("10.12", "Kümes hayvanları etlerinin işlenmesi ve saklanması", "10", "Gıda ürünlerinin imalatı"),
    ("10.13", "Et ve kümes hayvanları etlerinden üretilen ürünlerin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.20", "Balık, kabuklu deniz hayvanları ve yumuşakçaların işlenmesi ve saklanması", "10", "Gıda ürünlerinin imalatı"),
    ("10.31", "Patatesin işlenmesi ve saklanması", "10", "Gıda ürünlerinin imalatı"),
    ("10.32", "Sebze ve meyve suyu imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.39", "Başka yerde sınıflandırılmamış meyve ve sebzelerin işlenmesi ve saklanması", "10", "Gıda ürünlerinin imalatı"),
    ("10.41", "Sıvı ve katı yağ imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.42", "Margarin ve benzeri yenilebilir katı yağların imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.51", "Süthane işletmeciliği ve peynir imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.52", "Dondurma imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.61", "Öğütülmüş hububat ve sebze ürünleri imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.62", "Nişasta ve nişastalı ürünlerin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.71", "Ekmek, taze pastane ürünleri ve taze kek imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.72", "Peksimet ve bisküvi imalatı; dayanıklı pastane ürünleri ve dayanıklı kek imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.73", "Makarna, şehriye, kuskus ve benzeri unlu mamullerin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.81", "Şeker imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.82", "Kakao, çikolata ve şekerleme imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.83", "Kahve ve çayın işlenmesi", "10", "Gıda ürünlerinin imalatı"),
    ("10.84", "Baharat, sos, sirke ve diğer çeşni maddelerinin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.85", "Hazır yemeklerin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("10.89", "Başka yerde sınıflandırılmamış diğer gıda maddelerinin imalatı", "10", "Gıda ürünlerinin imalatı"),
    ("11.01", "Alkollü içeceklerin damıtılması, arıtılması ve harmanlanması", "11", "İçeceklerin imalatı"),
    ("11.02", "Üzümden şarap imalatı", "11", "İçeceklerin imalatı"),
    ("11.05", "Bira imalatı", "11", "İçeceklerin imalatı"),
    ("11.07", "Alkolsüz içeceklerin imalatı; maden sularının ve diğer şişelenmiş suların üretimi", "11", "İçeceklerin imalatı"),
    ("12.00", "Tütün ürünleri imalatı", "12", "Tütün ürünleri imalatı"),
    ("13.10", "Tekstil elyafının hazırlanması ve bükülmesi", "13", "Tekstil ürünlerinin imalatı"),
    ("13.20", "Dokuma", "13", "Tekstil ürünlerinin imalatı"),
    ("13.30", "Tekstil ürünlerinin bitirilmesi", "13", "Tekstil ürünlerinin imalatı"),
    ("13.91", "Örgü (triko) veya tığ işi (kroşe) kumaşların imalatı", "13", "Tekstil ürünlerinin imalatı"),
    ("13.92", "Giyim eşyası dışındaki tamamlanmış tekstil ürünlerinin imalatı", "13", "Tekstil ürünlerinin imalatı"),
    ("13.93", "Halı ve kilim imalatı", "13", "Tekstil ürünlerinin imalatı"),
    ("13.95", "Dokusuz kumaşların ve dokusuz kumaştan yapılan ürünlerin imalatı, giyim eşyası hariç", "13", "Tekstil ürünlerinin imalatı"),
    ("13.96", "Diğer teknik ve endüstriyel tekstillerin imalatı", "13", "Tekstil ürünlerinin imalatı"),
    ("13.99", "Başka yerde sınıflandırılmamış diğer tekstillerin imalatı", "13", "Tekstil ürünlerinin imalatı"),
    ("14.13", "Diğer dış giyim eşyaları imalatı", "14", "Giyim eşyalarının imalatı"),
    ("15.11", "Derinin tabaklanması ve işlenmesi; kürkün işlenmesi ve boyanması", "15", "Deri ve ilgili ürünlerin imalatı"),
    ("15.12", "Bavul, el çantası ve benzerleri ile saraçlık ve koşum takımı imalatı (deri giyim eşyası hariç)", "15", "Deri ve ilgili ürünlerin imalatı"),
    ("15.20", "Ayakkabı, bot, terlik vb. imalatı", "15", "Deri ve ilgili ürünlerin imalatı"),
    ("16.10", "Ağaçların biçilmesi ve planyalanması", "16", "Ağaç, ağaç ürünleri ve mantar ürünleri imalatı (mobilya hariç)"),
    ("16.21", "Ahşap kaplama paneli ve ağaç esaslı panel imalatı", "16", "Ağaç, ağaç ürünleri ve mantar ürünleri imalatı (mobilya hariç)"),
    ("16.22", "Birleştirilmiş parke yer döşemelerinin imalatı", "16", "Ağaç, ağaç ürünleri ve mantar ürünleri imalatı (mobilya hariç)"),
    ("16.23", "Diğer bina doğramacılığı ve marangozluk ürünlerinin imalatı", "16", "Ağaç, ağaç ürünleri ve mantar ürünleri imalatı (mobilya hariç)"),
    ("16.29", "Diğer ağaç ürünleri imalatı; mantardan, saz, saman ve benzeri örme malzemelerinden yapılmış ürünlerin imalatı", "16", "Ağaç, ağaç ürünleri ve mantar ürünleri imalatı (mobilya hariç)"),
    ("17.11", "Kağıt hamuru imalatı", "17", "Kağıt ve kağıt ürünlerinin imalatı"),
    ("17.12", "Kağıt ve mukavva imalatı", "17", "Kağıt ve kağıt ürünlerinin imalatı"),
    ("17.22", "Kağıttan yapılan ev eşyası, sıhhi malzemeler ve tuvalet malzemeleri imalatı", "17", "Kağıt ve kağıt ürünlerinin imalatı"),
    ("19.20", "Rafine edilmiş petrol ürünleri imalatı", "19", "Kok kömürü ve rafine edilmiş petrol ürünleri imalatı"),
    ("20.11", "Sanayi gazları imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.12", "Boya maddeleri ve pigment imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.13", "Diğer inorganik temel kimyasal maddelerin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.14", "Diğer organik temel kimyasalların imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.15", "Kimyasal gübre ve azot bileşiklerinin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.16", "Birincil formda plastik hammaddelerin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.17", "Birincil formda sentetik kauçuk imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.20", "Haşere ilaçları ve diğer zirai-kimyasal ürünlerin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.30", "Boya, vernik ve benzeri kaplayıcı maddeler ile matbaa mürekkebi ve macun imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.41", "Sabun ve deterjan ile temizlik ve parlatıcı maddeler imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.42", "Parfümlerin, kozmetiklerin ve kişisel bakım ürünlerinin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.59", "Başka yerde sınıflandırılmamış diğer kimyasal ürünlerin imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("20.60", "Suni veya sentetik elyaf imalatı", "20", "Kimyasalların ve kimyasal ürünlerin imalatı"),
    ("21.20", "Eczacılığa ilişkin ilaçların imalatı", "21", "Temel eczacılık ürünlerinin ve eczacılığa ilişkin malzemelerin imalatı"),
    ("22.11", "İç ve dış lastik imalatı; lastiğe sırt geçirilmesi ve yeniden işlenmesi", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("22.19", "Diğer kauçuk ürünleri imalatı", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("22.21", "Plastik tabaka, levha, tüp ve profil imalatı", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("22.22", "Plastik torba, çanta, poşet, çuval, kutu, damacana, şişe, makara vb. paketleme malzemelerinin imalatı", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("22.23", "Plastik inşaat malzemesi imalatı", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("22.29", "Diğer plastik ürünlerin imalatı", "22", "Kauçuk ve plastik ürünlerin imalatı"),
    ("23.11", "Düz cam imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.13", "Çukur cam imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.14", "Cam elyafı imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.20", "Ateşe dayanıklı (refrakter) ürünlerin imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.31", "Seramik karo ve kaldırım taşları imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.41", "Seramik ev ve süs eşyaları imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.42", "Seramik sıhhi ürünlerin imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.52", "Kireç ve alçı imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.61", "İnşaat amaçlı beton ürünlerin imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.62", "İnşaat amaçlı alçı ürünlerin imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.64", "Toz harç imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("23.99", "Başka yerde sınıflandırılmamış metalik olmayan diğer mineral ürünlerin imalatı", "23", "Diğer metalik olmayan mineral ürünlerin imalatı"),
    ("24.10", "Ana demir ve çelik ürünleri ile ferro alaşımların imalatı", "24", "Ana metal sanayii"),
    ("24.20", "Çelikten tüpler, borular, içi boş profiller ve benzeri bağlantı parçalarının imalatı", "24", "Ana metal sanayii"),
    ("24.31", "Barların soğuk çekilmesi", "24", "Ana metal sanayii"),
    ("24.32", "Dar şeritlerin soğuk haddelenmesi", "24", "Ana metal sanayii"),
    ("24.34", "Tellerin soğuk çekilmesi", "24", "Ana metal sanayii"),
    ("24.41", "Değerli metal üretimi", "24", "Ana metal sanayii"),
    ("24.42", "Alüminyum üretimi", "24", "Ana metal sanayii"),
    ("24.51", "Demir döküm", "24", "Ana metal sanayii"),
    ("24.52", "Çelik dökümü", "24", "Ana metal sanayii"),
    ("24.53", "Hafif metallerin dökümü", "24", "Ana metal sanayii"),
    ("24.54", "Diğer demir dışı metallerin dökümü", "24", "Ana metal sanayii"),
    ("25.12", "Metalden kapı ve pencere imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.21", "Merkezi ısıtma radyatörleri (elektrikli radyatörler hariç) ve sıcak su kazanları (boylerleri) imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.30", "Buhar jeneratörü imalatı, merkezi ısıtma sıcak su kazanları (boylerleri) hariç", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.50", "Metallerin dövülmesi, preslenmesi, baskılanması ve yuvarlanması; toz metalürjisi", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.61", "Metallerin işlenmesi ve kaplanması", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.62", "Metallerin makinede işlenmesi ve şekil verilmesi", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.71", "Çatal-bıçak takımları ve diğer kesici aletlerin imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.73", "El aletleri, takım tezgahı uçları, testere ağızları vb. imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.92", "Metalden hafif paketleme malzemeleri imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.93", "Tel ürünleri, zincir ve yayların imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.94", "Bağlantı malzemelerinin ve vida makinesi ürünlerinin imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("25.99", "Başka yerde sınıflandırılmamış diğer fabrikasyon metal ürünlerin imalatı", "25", "Fabrikasyon metal ürünleri imalatı (makine ve teçhizat hariç)"),
    ("26.40", "Tüketici elektroniği ürünlerinin imalatı", "26", "Bilgisayarların, elektronik ve optik ürünlerin imalatı"),
    ("26.51", "Ölçme, test ve seyrüsefer amaçlı alet ve cihazların imalatı", "26", "Bilgisayarların, elektronik ve optik ürünlerin imalatı"),
    ("27.11", "Elektrik motorlarının, jeneratörlerin ve transformatörlerin imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.12", "Elektrik dağıtım ve kontrol cihazları imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.20", "Akümülatör ve pil imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.31", "Fiber optik kabloların imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.32", "Diğer elektronik ve elektrik telleri ve kablolarının imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.40", "Elektrikli aydınlatma ekipmanlarının imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("27.51", "Elektrikli ev aletlerinin imalatı", "27", "Elektrikli teçhizat imalatı"),
    ("28.12", "Akışkan gücü ile çalışan ekipmanların imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.14", "Diğer musluk ve valf/vana imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.15", "Rulman, dişli/dişli takımı, şanzıman ve tahrik elemanlarının imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.22", "Kaldırma ve taşıma ekipmanları imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.25", "Soğutma ve havalandırma donanımlarının imalatı, evde kullanılanlar hariç", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.92", "Maden, taş ocağı ve inşaat makineleri imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.94", "Tekstil, giyim eşyası ve deri üretiminde kullanılan makinelerin imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("28.99", "Başka yerde sınıflandırılmamış diğer özel amaçlı makinelerin imalatı", "28", "Başka yerde sınıflandırılmamış makine ve ekipman imalatı"),
    ("29.10", "Motorlu kara taşıtlarının imalatı", "29", "Motorlu kara taşıtı, treyler (römork) ve yarı treyler (yarı römork) imalatı"),
    ("29.20", "Motorlu kara taşıtları karoseri (kaporta) imalatı; treyler (römork) ve yarı treyler (yarı römork) imalatı", "29", "Motorlu kara taşıtı, treyler (römork) ve yarı treyler (yarı römork) imalatı"),
    ("29.32", "Motorlu kara taşıtları için diğer parça ve aksesuarların imalatı", "29", "Motorlu kara taşıtı, treyler (römork) ve yarı treyler (yarı römork) imalatı"),
    ("30.11", "Gemilerin ve yüzen yapıların inşası", "30", "Diğer ulaşım araçlarının imalatı"),
    ("30.20", "Demir yolu lokomotifleri ve vagonlarının imalatı", "30", "Diğer ulaşım araçlarının imalatı"),
    ("32.12", "Mücevher ve benzeri eşyaların imalatı", "32", "Diğer imalatlar"),
    ("32.40", "Oyun ve oyuncak imalatı", "32", "Diğer imalatlar"),
    ("33.16", "Hava taşıtlarının ve uzay araçlarının bakım ve onarımı", "33", "Makine ve ekipmanların kurulumu ve onarımı"),
    ("33.20", "Sanayi makine ve ekipmanlarının kurulumu", "33", "Makine ve ekipmanların kurulumu ve onarımı"),
    ("35.11", "Elektrik enerjisi üretimi", "35", "Elektrik, gaz, buhar ve havalandırma sistemi üretim ve dağıtımı"),
    ("35.30", "Buhar ve iklimlendirme temini", "35", "Elektrik, gaz, buhar ve havalandırma sistemi üretim ve dağıtımı"),
    ("38.32", "Tasnif edilmiş materyallerin geri kazanımı", "38", "Atığın toplanması, ıslahı ve bertarafı faaliyetleri; maddelerin geri kazanımı"),
    ("42.12", "Demir yolları ve metroların inşaatı", "42", "Bina dışı yapıların inşaatı"),
]

FIRST_DATA_ROW = 10
LAST_DATA_ROW = 59
NACE_FIRST = 2
NACE_LAST = 1 + len(EK2_NACE)

NAVY = "1B365D"
NAVY_DARK = "0F2340"
GOLD = "C4A35A"
YELLOW = "FFF2CC"
LIGHT_BLUE = "D6E3F0"
GRAY = "F3F6F9"
RESULT_GRAY = "EEF2F6"
WHITE = "FFFFFF"
RED = "C00000"
GREEN = "548235"
SOFT_RED = "F4CCCC"
SOFT_GREEN = "D9EAD3"
thin = Border(
    left=Side(style="thin", color="B7C4D4"),
    right=Side(style="thin", color="B7C4D4"),
    top=Side(style="thin", color="B7C4D4"),
    bottom=Side(style="thin", color="B7C4D4"),
)
thick_bottom = Border(
    left=Side(style="thin", color="B7C4D4"),
    right=Side(style="thin", color="B7C4D4"),
    top=Side(style="thin", color="B7C4D4"),
    bottom=Side(style="medium", color=NAVY),
)


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def apply_common(cell, *, font=None, fill_color=None, align=None, border=thin, wrap=True):
    cell.font = font or Font(name="Calibri", size=11, color="1F2933")
    if fill_color:
        cell.fill = fill(fill_color)
    cell.alignment = align or Alignment(vertical="center", wrap_text=wrap)
    cell.border = border


def digits_formula(cell_ref: str) -> str:
    return (
        f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({cell_ref}),".",""),",",""),"-","")," ",""),";","")'
    )


def normalized_nace_formula(row: int) -> str:
    digits = digits_formula(f"B{row}")
    return (
        f'IF(B{row}="","",'
        f'IF(LEN({digits})<4,"GEÇERSİZ",'
        f'LEFT({digits},2)&"."&MID({digits},3,2)))'
    )


def nace_match_formula(row: int) -> str:
    return (
        f'IF(OR(B{row}="",I{row}="",I{row}="GEÇERSİZ"),"",'
        f'IF(COUNTIF(\'Ek-2 NACE Listesi\'!$B${NACE_FIRST}:$B${NACE_LAST},I{row})>0,"Evet","Hayır"))'
    )


def employee_formula(row: int) -> str:
    return (
        f'IF(C{row}="","",'
        f'IF(AND(ISNUMBER(C{row}),C{row}>=50),"Evet","Hayır"))'
    )


def scope_formula(row: int) -> str:
    return (
        f'IF(OR(B{row}="",C{row}=""),"",'
        f'IF(AND(D{row}="Evet",E{row}="Evet"),"KAPSAMDA","KAPSAM DIŞI"))'
    )


def reason_formula(row: int) -> str:
    return (
        f'IF(OR(B{row}="",C{row}=""),"",'
        f'IF(I{row}="GEÇERSİZ","NACE kodunu 4 haneli yazın (ör. 10.11)",'
        f'IF(AND(D{row}="Evet",E{row}="Evet"),"NACE kodu Ek-2 listesinde ve çalışan sayısı 50 ve üzeri",'
        f'IF(AND(D{row}="Evet",E{row}="Hayır"),"NACE kodu listede ancak çalışan sayısı 50\'den az",'
        f'IF(AND(D{row}="Hayır",E{row}="Evet"),"Çalışan sayısı yeterli ancak NACE kodu Ek-2 listesinde yok",'
        f'"NACE kodu Ek-2 listesinde yok ve çalışan sayısı 50\'den az")))))'
    )


def activity_formula(row: int) -> str:
    return (
        f'IF(OR(B{row}="",I{row}="",I{row}="GEÇERSİZ"),"",'
        f'IFERROR(VLOOKUP(I{row},\'Ek-2 NACE Listesi\'!$B${NACE_FIRST}:$C${NACE_LAST},2,FALSE),"Listede bulunamadı"))'
    )


def build_nace_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Ek-2 NACE Listesi")
    ws.sheet_properties.tabColor = NAVY

    headers = ["Sıra", "NACE Kodu", "Faaliyet Açıklaması", "Ana Faaliyet Kodu", "Ana Faaliyet Açıklaması"]
    widths = [10, 16, 78, 20, 62]
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, col, header)
        apply_common(
            cell,
            font=Font(name="Calibri", size=11, bold=True, color=WHITE),
            fill_color=NAVY,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True),
        )
        ws.column_dimensions[get_column_letter(col)].width = width

    for idx, (code, desc, ana_kod, ana_desc) in enumerate(EK2_NACE, start=1):
        row = idx + 1
        values = [idx, code, desc, ana_kod, ana_desc]
        bg = WHITE if idx % 2 else GRAY
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            if col == 2:
                cell.number_format = "@"
            apply_common(
                cell,
                font=Font(name="Calibri", size=11, bold=(col == 2), color=NAVY if col == 2 else "1F2933"),
                fill_color=bg,
                align=Alignment(
                    horizontal="center" if col in (1, 2, 4) else "left",
                    vertical="center",
                    wrap_text=True,
                ),
            )
        ws.row_dimensions[row].height = 28

    ws.auto_filter.ref = f"A1:E{NACE_LAST}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False
    ws.oddHeader.left.text = "Su Verimliliği Yönetmeliği Ek-2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def build_help_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Nasıl Kullanılır")
    ws.sheet_properties.tabColor = GOLD
    ws.column_dimensions["A"].width = 118

    title = ws["A1"]
    title.value = "Su Verimliliği Yönetmeliği — Kapsam Kontrol Tablosu Kullanım Kılavuzu"
    apply_common(
        title,
        font=Font(name="Calibri", size=16, bold=True, color=WHITE),
        fill_color=NAVY,
        align=Alignment(vertical="center", wrap_text=True),
        border=None,
    )
    ws.row_dimensions[1].height = 36

    blocks = [
        ("Kural", "Firma KAPSAMDA sayılır ancak şu iki şart birlikte sağlanır:\n1) NACE kodu, 27.12.2024 tarihli Su Verimliliği Yönetmeliği Ek-2 listesinde yer alır.\n2) Çalışan sayısı 50 ve üzeridir (yönetmelikteki “50 ve üzeri çalışan” ifadesi).\nAksi halde sonuç KAPSAM DIŞI olur."),
        ("Nasıl doldurulur?", "Sadece “Firma Kontrolü” sayfasındaki sarı hücrelere yazın:\n• Firma Adı\n• NACE Kodu (ör. 10.11 veya 10.11.01 — tablo ilk dört haneyi kullanır)\n• Çalışan Sayısı\nKapsam Durumu, gerekçe ve faaliyet açıklaması otomatik dolar. Formül hücrelerini değiştirmeyin."),
        ("Renkler", "Kırmızı / KAPSAMDA = NACE listede + çalışan sayısı ≥ 50.\nYeşil / KAPSAM DIŞI = NACE listede değil veya çalışan sayısı 50’nin altında."),
        ("Kaynak", "Resmî Gazete: 27 Aralık 2024, Sayı 32765 — Su Verimliliği Yönetmeliği, Ek-2 NACE Kodu Listesi (148 kod).\nEndüstriyel tesisler için sistem kurulum süresi yönetmelik yayımından itibaren 18 aydır."),
        ("Not", "Organize sanayi bölgeleri, serbest bölgeler ve endüstri bölgeleri NACE/çalışan sayısından bağımsız olarak ayrıca yükümlüdür. Bu tablo münferit endüstriyel işletmeler içindir.\n13 Mart 2025 tarihli kılavuz güncellemesinde bazı Ek-2 kodları gönüllü statüsüne alınmıştır. Bu dosya yönetmeliğin Ek-2 listesine ve sizin belirttiğiniz iki şarta göre çalışır."),
    ]

    row = 3
    for heading, body in blocks:
        h = ws.cell(row, 1, heading)
        apply_common(
            h,
            font=Font(name="Calibri", size=12, bold=True, color=NAVY),
            fill_color=LIGHT_BLUE,
            border=None,
        )
        ws.row_dimensions[row].height = 22
        row += 1
        b = ws.cell(row, 1, body)
        apply_common(
            b,
            font=Font(name="Calibri", size=11, color="1F2933"),
            fill_color=WHITE,
            align=Alignment(vertical="top", wrap_text=True),
            border=None,
        )
        ws.row_dimensions[row].height = 78 if heading != "Not" else 92
        row += 2

    ws.sheet_view.showGridLines = False


def build_control_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Firma Kontrolü", 0)
    ws.sheet_properties.tabColor = RED
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"

    ws.merge_cells("A1:I2")
    title = ws["A1"]
    title.value = "SU VERİMLİLİĞİ YÖNETMELİĞİ  •  FİRMA KAPSAM KONTROLÜ"
    apply_common(
        title,
        font=Font(name="Calibri", size=20, bold=True, color=WHITE),
        fill_color=NAVY,
        align=Alignment(horizontal="center", vertical="center"),
        border=None,
    )
    ws["B1"].fill = fill(NAVY)
    for col in range(1, 10):
        for r in (1, 2):
            ws.cell(r, col).fill = fill(NAVY)
            ws.cell(r, col).border = Border()
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:I3")
    subtitle = ws["A3"]
    subtitle.value = "Kaynak: 27.12.2024 tarihli Su Verimliliği Yönetmeliği Ek-2 NACE listesi  |  Kural: NACE kodu listede + çalışan sayısı 50 ve üzeri  →  KAPSAMDA"
    apply_common(
        subtitle,
        font=Font(name="Calibri", size=10, italic=True, color=NAVY),
        fill_color=LIGHT_BLUE,
        align=Alignment(horizontal="left", vertical="center"),
        border=None,
    )
    for col in range(1, 10):
        ws.cell(3, col).fill = fill(LIGHT_BLUE)
        ws.cell(3, col).border = Border()
    ws.row_dimensions[3].height = 22

    ws.merge_cells("A4:C4")
    ws.merge_cells("D4:F4")
    ws.merge_cells("G4:I4")
    legend = [
        (4, 1, "Sarı hücrelere yazın: firma adı, NACE, çalışan sayısı", YELLOW, NAVY),
        (4, 4, "KAPSAMDA  =  kırmızı", RED, WHITE),
        (4, 7, "KAPSAM DIŞI  =  yeşil", GREEN, WHITE),
    ]
    for row, col, text, bg, fg in legend:
        cell = ws.cell(row, col, text)
        apply_common(
            cell,
            font=Font(name="Calibri", size=10, bold=True, color=fg),
            fill_color=bg,
            align=Alignment(horizontal="center", vertical="center"),
            border=None,
        )
    for col in range(1, 4):
        ws.cell(4, col).fill = fill(YELLOW)
        ws.cell(4, col).border = Border()
    for col in range(4, 7):
        ws.cell(4, col).fill = fill(RED)
        ws.cell(4, col).border = Border()
        ws.cell(4, col).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    for col in range(7, 10):
        ws.cell(4, col).fill = fill(GREEN)
        ws.cell(4, col).border = Border()
        ws.cell(4, col).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    ws.row_dimensions[4].height = 22

    ws.merge_cells("A5:I5")
    note = ws["A5"]
    note.value = (
        "İlk üç satır örnektir; silebilir veya üzerine yazabilirsiniz. "
        "NACE kodunu 10.11, 10,11 veya 10.11.01 biçiminde yazabilirsiniz. "
        "6 haneli kod girilirse ilk dört hane (10.11) kontrol edilir."
    )
    apply_common(
        note,
        font=Font(name="Calibri", size=10, color="334155"),
        fill_color=GRAY,
        align=Alignment(horizontal="left", vertical="center", wrap_text=True),
        border=None,
    )
    for col in range(1, 10):
        ws.cell(5, col).fill = fill(GRAY)
        ws.cell(5, col).border = Border()
    ws.row_dimensions[5].height = 32

    ws.merge_cells("A6:C6")
    ws.merge_cells("D6:F6")
    summary_labels = [
        (6, 1, "Özet", NAVY, WHITE, True),
        (6, 4, "KAPSAMDA firma sayısı", RED, WHITE, True),
        (6, 7, "KAPSAM DIŞI firma sayısı", GREEN, WHITE, True),
    ]
    for row, col, text, bg, fg, bold in summary_labels:
        cell = ws.cell(row, col, text)
        apply_common(
            cell,
            font=Font(name="Calibri", size=10, bold=bold, color=fg),
            fill_color=bg,
            align=Alignment(horizontal="center", vertical="center"),
            border=None,
        )
    for col in range(1, 4):
        ws.cell(6, col).fill = fill(NAVY)
        ws.cell(6, col).font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        ws.cell(6, col).border = Border()
    for col in range(4, 7):
        ws.cell(6, col).fill = fill(RED)
        ws.cell(6, col).border = Border()
    for col in range(7, 10):
        ws.cell(6, col).fill = fill(GREEN)
        ws.cell(6, col).border = Border()

    ws.merge_cells("A7:C7")
    ws.merge_cells("D7:F7")
    ws.merge_cells("G7:I7")
    ws["A7"].value = f"Ek-2 listesindeki NACE kodu adedi: {len(EK2_NACE)}"
    ws["D7"].value = f'=COUNTIF(F{FIRST_DATA_ROW}:F{LAST_DATA_ROW},"KAPSAMDA")'
    ws["G7"].value = f'=COUNTIF(F{FIRST_DATA_ROW}:F{LAST_DATA_ROW},"KAPSAM DIŞI")'
    apply_common(
        ws["A7"],
        font=Font(name="Calibri", size=12, bold=True, color=NAVY),
        fill_color=WHITE,
        align=Alignment(horizontal="center", vertical="center"),
        border=None,
    )
    apply_common(
        ws["D7"],
        font=Font(name="Calibri", size=16, bold=True, color=RED),
        fill_color=SOFT_RED,
        align=Alignment(horizontal="center", vertical="center"),
        border=None,
    )
    apply_common(
        ws["G7"],
        font=Font(name="Calibri", size=16, bold=True, color=GREEN),
        fill_color=SOFT_GREEN,
        align=Alignment(horizontal="center", vertical="center"),
        border=None,
    )
    for col in range(1, 4):
        ws.cell(7, col).fill = fill(WHITE)
        ws.cell(7, col).border = Border()
    for col in range(4, 7):
        ws.cell(7, col).fill = fill(SOFT_RED)
        ws.cell(7, col).border = Border()
    for col in range(7, 10):
        ws.cell(7, col).fill = fill(SOFT_GREEN)
        ws.cell(7, col).border = Border()
    ws.row_dimensions[6].height = 18
    ws.row_dimensions[7].height = 28
    ws.row_dimensions[8].height = 8

    headers = [
        "Firma Adı",
        "NACE Kodu",
        "Çalışan Sayısı",
        "NACE Listede mi?",
        "Çalışan ≥ 50 mi?",
        "Kapsam Durumu",
        "Gerekçe",
        "Ek-2 Faaliyet Açıklaması",
        "Normalize NACE",
    ]
    widths = [34, 16, 16, 18, 18, 18, 52, 62, 16]
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(9, col, header)
        apply_common(
            cell,
            font=Font(name="Calibri", size=11, bold=True, color=WHITE),
            fill_color=NAVY_DARK if col >= 4 else NAVY,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True),
            border=thick_bottom,
        )
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[9].height = 32

    examples = [
        ("Örnek Gıda A.Ş.", "10.11", 120),
        ("Örnek Market Ltd.", "47.11", 200),
        ("Örnek Küçük Tesis", "10.11", 30),
    ]

    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        example_idx = row - FIRST_DATA_ROW
        if example_idx < len(examples):
            name, nace, employees = examples[example_idx]
        else:
            name, nace, employees = "", "", ""

        input_bg = YELLOW
        result_bg = RESULT_GRAY if row % 2 else WHITE

        ws.cell(row, 1, name)
        nace_cell = ws.cell(row, 2, nace)
        nace_cell.number_format = "@"
        ws.cell(row, 3, employees if employees != "" else None)
        ws.cell(row, 4, f"={nace_match_formula(row)}")
        ws.cell(row, 5, f"={employee_formula(row)}")
        ws.cell(row, 6, f"={scope_formula(row)}")
        ws.cell(row, 7, f"={reason_formula(row)}")
        ws.cell(row, 8, f"={activity_formula(row)}")
        ws.cell(row, 9, f"={normalized_nace_formula(row)}")

        for col in range(1, 10):
            cell = ws.cell(row, col)
            is_input = col <= 3
            apply_common(
                cell,
                font=Font(
                    name="Calibri",
                    size=11,
                    bold=(col == 6),
                    color=NAVY if col in (2, 6) else "1F2933",
                ),
                fill_color=input_bg if is_input else result_bg,
                align=Alignment(
                    horizontal="center" if col in (2, 3, 4, 5, 6, 9) else "left",
                    vertical="center",
                    wrap_text=True,
                ),
            )
            if col == 2:
                cell.number_format = "@"
            if col == 3:
                cell.number_format = "0"
            if col == 9:
                cell.number_format = "@"
        ws.row_dimensions[row].height = 26

    red_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    green_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    ws.conditional_formatting.add(
        f"F{FIRST_DATA_ROW}:F{LAST_DATA_ROW}",
        CellIsRule(operator="equal", formula=['"KAPSAMDA"'], fill=fill(RED), font=red_font),
    )
    ws.conditional_formatting.add(
        f"F{FIRST_DATA_ROW}:F{LAST_DATA_ROW}",
        CellIsRule(operator="equal", formula=['"KAPSAM DIŞI"'], fill=fill(GREEN), font=green_font),
    )
    ws.conditional_formatting.add(
        f"D{FIRST_DATA_ROW}:E{LAST_DATA_ROW}",
        CellIsRule(
            operator="equal",
            formula=['"Evet"'],
            fill=fill(SOFT_GREEN),
            font=Font(name="Calibri", size=11, bold=True, color=GREEN),
        ),
    )
    ws.conditional_formatting.add(
        f"D{FIRST_DATA_ROW}:E{LAST_DATA_ROW}",
        CellIsRule(
            operator="equal",
            formula=['"Hayır"'],
            fill=fill(SOFT_RED),
            font=Font(name="Calibri", size=11, bold=True, color=RED),
        ),
    )

    nace_dv = DataValidation(
        type="list",
        formula1=f"='Ek-2 NACE Listesi'!$B${NACE_FIRST}:$B${NACE_LAST}",
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
        showInputMessage=True,
        promptTitle="NACE kodu",
        prompt="Listeden seçebilir veya kendi kodunuzu yazabilirsiniz. 6 haneli kodlar 4 haneye indirilir.",
    )
    nace_dv.add(f"B{FIRST_DATA_ROW}:B{LAST_DATA_ROW}")
    ws.add_data_validation(nace_dv)

    emp_dv = DataValidation(
        type="whole",
        operator="greaterThanOrEqual",
        formula1="0",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Geçersiz sayı",
        error="Çalışan sayısı 0 veya daha büyük bir tam sayı olmalıdır.",
        promptTitle="Çalışan sayısı",
        prompt="SGK’daki güncel çalışan sayısını yazın. 50 ve üzeri şartı sağlar.",
        showInputMessage=True,
    )
    emp_dv.add(f"C{FIRST_DATA_ROW}:C{LAST_DATA_ROW}")
    ws.add_data_validation(emp_dv)

    ws.auto_filter.ref = f"A9:I{LAST_DATA_ROW}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:9"
    ws.oddHeader.left.text = "Su Verimliliği Yönetmeliği — Firma Kapsam Kontrolü"
    ws.oddFooter.right.text = "Ek-2 NACE listesi  |  Sayfa &P / &N"
    ws.sheet_view.zoomScale = 110


def create_workbook(path: Path = OUTPUT_PATH) -> Path:
    if len(EK2_NACE) != 148:
        raise RuntimeError(f"Ek-2 listesi 148 kod olmalı, {len(EK2_NACE)} bulundu.")
    codes = [item[0] for item in EK2_NACE]
    if len(set(codes)) != len(codes):
        raise RuntimeError("Ek-2 listesinde tekrarlayan NACE kodu var.")

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    build_nace_sheet(wb)
    build_help_sheet(wb)
    build_control_sheet(wb)

    wb.properties.title = "Su Verimliliği Yönetmeliği Firma Kapsam Kontrolü"
    wb.properties.creator = "Su Verimliliği Kapsam Tablosu"
    wb.properties.subject = "Ek-2 NACE listesi ve 50+ çalışan kuralı"
    wb["Firma Kontrolü"].sheet_view.tabSelected = True
    wb.active = wb["Firma Kontrolü"]

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _make_excel_compatible(path)
    return path


def _make_excel_compatible(path: Path) -> None:
    """Excel'in dosyayı onarım istemeden açması için koruma etiketlerini temizler."""
    with ZipFile(path, "r") as src:
        parts = {name: src.read(name) for name in src.namelist()}

    workbook = parts["xl/workbook.xml"].decode("utf-8")
    workbook = workbook.replace("<workbookProtection />", "")
    workbook = workbook.replace("<workbookProtection/>", "")
    parts["xl/workbook.xml"] = workbook.encode("utf-8")

    for name, data in list(parts.items()):
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            text = data.decode("utf-8")
            if "<sheetProtection" in text:
                start = text.index("<sheetProtection")
                end = text.index("/>", start) + 2
                text = text[:start] + text[end:]
                parts[name] = text.encode("utf-8")

    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as dest:
        for name, data in parts.items():
            dest.writestr(name, data)
    path.write_bytes(buffer.getvalue())


if __name__ == "__main__":
    output = create_workbook()
    print(f"Excel oluşturuldu: {output}")
    print(f"Ek-2 NACE adedi: {len(EK2_NACE)}")
