#!/usr/bin/env python3
"""Boru üretim prosesi iş akım şeması — emisyon bacaları işaretli."""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

OUT_DIR = Path(__file__).resolve().parent

NAVY = "#1F4E79"
GREEN = "#1E7A46"
ORANGE = "#E67E22"
CHIMNEY_FILL = "#FFF3E0"
CHIMNEY_STROKE = "#E65100"
CHIMNEY_TEXT = "#BF360C"

STEPS = [
    {
        "id": "s1",
        "title": "1. Hammadde girdi kontrol ve stoklama",
        "sub": "Cam elyaf, kum, reçine",
        "kind": "main",
    },
    {
        "id": "s2",
        "title": "2. Winder'da boru üretimi",
        "sub": "Günlük tank / raf — döner kalıba sarım",
        "kind": "main",
        "baca": "Stiren Emiş Bacası 1",
    },
    {
        "id": "s3",
        "title": "3. Boru boy kesme",
        "sub": "",
        "kind": "main",
    },
    {
        "id": "sodd",
        "title": "ODD tamir ünitesi — ihtiyaç halinde",
        "sub": "Hata yoksa bu adım atlanır",
        "kind": "optional",
        "baca": "ODD Tamir Proses Bacası",
    },
    {
        "id": "s4",
        "title": "4. Boru ucu kalibrasyon",
        "sub": "",
        "kind": "main",
        "baca": "Kalibrasyon Proses Bacası",
    },
    {
        "id": "s5",
        "title": "5. Hidrotest",
        "sub": "",
        "kind": "main",
    },
    {
        "id": "s6",
        "title": "6. Coupling kanal açma",
        "sub": "",
        "kind": "main",
        "baca": "Kaplin Kanal Açma Proses Bacası",
    },
    {
        "id": "s7",
        "title": "7. Stopper ve conta montajı",
        "sub": "",
        "kind": "main",
    },
    {
        "id": "s8",
        "title": "8. Stok sahasına sevk",
        "sub": "İstifleme — sevkiyat — müşteri",
        "kind": "end",
    },
]


def _html_value(title: str, sub: str) -> str:
    title_e = escape(title)
    if not sub:
        return title_e
    return f"{title_e}&lt;br&gt;&lt;font style=&quot;font-size:11px;font-weight:normal;&quot;&gt;{escape(sub)}&lt;/font&gt;"


def write_drawio(path: Path) -> None:
    box_w, box_h = 390, 78
    baca_w, baca_h = 280, 52
    x0, y0 = 80, 90
    gap = 28
    baca_x = x0 + box_w + 70
    page_w = 900
    page_h = 1180

    cells = [
        '        <mxCell id="0"/>',
        '        <mxCell id="1" parent="0"/>',
        '        <mxCell id="proses" value="Proses" parent="0"/>',
        '        <mxCell id="bacalar" value="Bacalar" parent="0"/>',
        '        <mxCell id="yazi" value="Yazilar" parent="0"/>',
    ]

    cells.append(
        '        <mxCell id="baslik" value="Boru Üretim Tesisi — Proses İş Akım Şeması" '
        'style="text;html=1;align=center;fontSize=20;fontStyle=1;fontFamily=Arial;'
        'fontColor=#1F4E79;" vertex="1" parent="yazi">'
        f'\n          <mxGeometry x="40" y="18" width="820" height="32" as="geometry"/>\n'
        "        </mxCell>"
    )
    cells.append(
        '        <mxCell id="altbaslik" value="Emisyon bacaları ilgili proses adımlarının yanında ok ile gösterilmiştir." '
        'style="text;html=1;align=center;fontSize=12;fontFamily=Arial;fontColor=#555555;" '
        'vertex="1" parent="yazi">'
        f'\n          <mxGeometry x="40" y="50" width="820" height="24" as="geometry"/>\n'
        "        </mxCell>"
    )

    ys = []
    y = y0
    for step in STEPS:
        ys.append(y)
        kind = step["kind"]
        if kind == "optional":
            style = (
                "rounded=1;whiteSpace=wrap;html=1;arcSize=12;dashed=1;dashPattern=8 6;"
                f"fillColor=#ffffff;strokeColor={ORANGE};fontColor={ORANGE};"
                "fontStyle=1;fontSize=14;fontFamily=Arial;strokeWidth=2;"
            )
        elif kind == "end":
            style = (
                "rounded=1;whiteSpace=wrap;html=1;arcSize=12;"
                f"fillColor={GREEN};strokeColor={GREEN};fontColor=#FFFFFF;"
                "fontStyle=1;fontSize=14;fontFamily=Arial;strokeWidth=1;"
            )
        else:
            style = (
                "rounded=1;whiteSpace=wrap;html=1;arcSize=12;"
                f"fillColor={NAVY};strokeColor={NAVY};fontColor=#FFFFFF;"
                "fontStyle=1;fontSize=14;fontFamily=Arial;strokeWidth=1;"
            )
        value = _html_value(step["title"], step["sub"])
        cells.append(
            f'        <mxCell id="{step["id"]}" value="{value}" style="{style}" '
            f'vertex="1" parent="proses">\n'
            f'          <mxGeometry x="{x0}" y="{y}" width="{box_w}" height="{box_h}" as="geometry"/>\n'
            "        </mxCell>"
        )
        if step.get("baca"):
            baca_style = (
                "rounded=1;whiteSpace=wrap;html=1;arcSize=10;"
                f"fillColor={CHIMNEY_FILL};strokeColor={CHIMNEY_STROKE};fontColor={CHIMNEY_TEXT};"
                "fontStyle=1;fontSize=13;fontFamily=Arial;strokeWidth=2;"
            )
            baca_id = f'{step["id"]}-baca'
            baca_y = y + (box_h - baca_h) / 2
            cells.append(
                f'        <mxCell id="{baca_id}" value="{escape(step["baca"])}" style="{baca_style}" '
                f'vertex="1" parent="bacalar">\n'
                f'          <mxGeometry x="{baca_x}" y="{baca_y}" width="{baca_w}" '
                f'height="{baca_h}" as="geometry"/>\n'
                "        </mxCell>"
            )
            cells.append(
                f'        <mxCell id="{step["id"]}-ok" style="endArrow=block;endFill=1;html=1;'
                f"strokeColor={CHIMNEY_STROKE};strokeWidth=2.2;exitX=1;exitY=0.5;exitDx=0;exitDy=0;"
                f'entryX=0;entryY=0.5;entryDx=0;entryDy=0;" edge="1" parent="bacalar" '
                f'source="{step["id"]}" target="{baca_id}">\n'
                '          <mxGeometry relative="1" as="geometry"/>\n'
                "        </mxCell>"
            )
        y += box_h + gap

    for i in range(len(STEPS) - 1):
        src = STEPS[i]["id"]
        dst = STEPS[i + 1]["id"]
        dashed = STEPS[i + 1]["kind"] == "optional" or STEPS[i]["kind"] == "optional"
        if dashed and STEPS[i + 1]["kind"] == "optional":
            edge_style = (
                f"endArrow=block;endFill=1;html=1;dashed=1;dashPattern=8 6;"
                f"strokeColor={ORANGE};strokeWidth=2.2;"
            )
        else:
            edge_style = (
                f"endArrow=block;endFill=1;html=1;strokeColor={NAVY};strokeWidth=2.2;"
            )
        cells.append(
            f'        <mxCell id="e-{src}-{dst}" style="{edge_style}exitX=0.5;exitY=1;'
            f'exitDx=0;exitDy=0;entryX=0.5;entryY=0;entryDx=0;entryDy=0;" edge="1" '
            f'parent="proses" source="{src}" target="{dst}">\n'
            '          <mxGeometry relative="1" as="geometry"/>\n'
            "        </mxCell>"
        )

    cells.append(
        '        <mxCell id="lejant" value="'
        + escape(
            "Lejant\n■ Proses adımı\n□ ODD tamir (opsiyonel)\n■ Sevk\n■ Emisyon bacası"
        ).replace("\n", "&lt;br&gt;")
        + '" style="rounded=0;whiteSpace=wrap;html=1;align=left;verticalAlign=top;'
        "fillColor=#f7f7f7;strokeColor=#cccccc;fontSize=11;fontFamily=Arial;spacing=8;\" "
        'vertex="1" parent="yazi">\n'
        '          <mxGeometry x="80" y="1088" width="220" height="78" as="geometry"/>\n'
        "        </mxCell>"
    )

    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<mxfile host="app.diagrams.net" agent="GES-Metraj-Pro" version="22.1.0">
  <diagram id="boru-is-akimi" name="Boru Üretim İş Akım Şeması">
    <mxGraphModel dx="1200" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="{page_w}" pageHeight="{page_h}" math="0" shadow="0">
      <root>
{chr(10).join(cells)}
      </root>
    </mxGraphModel>
  </diagram>
</mxfile>
"""
    path.write_text(xml, encoding="utf-8")


def write_png(path: Path) -> None:
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

    fig, ax = plt.subplots(figsize=(10.2, 14.4), dpi=160)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    ax.set_xlim(0, 10.2)
    ax.set_ylim(0, 14.4)
    ax.invert_yaxis()
    ax.axis("off")

    ax.text(
        5.1,
        0.28,
        "Boru Üretim Tesisi — Proses İş Akım Şeması",
        ha="center",
        va="top",
        fontsize=16,
        fontweight="bold",
        color=NAVY,
    )
    ax.text(
        5.1,
        0.62,
        "Emisyon bacaları ilgili proses adımlarının yanında ok ile gösterilmiştir.",
        ha="center",
        va="top",
        fontsize=9.5,
        color="#555555",
    )

    box_w, box_h = 4.35, 0.95
    x0 = 0.55
    baca_x = 6.15
    baca_w, baca_h = 3.55, 0.62
    y = 1.05
    gap = 0.38
    centers = []

    def wrap_title(text: str, width: int = 42) -> str:
        return text

    for step in STEPS:
        kind = step["kind"]
        if kind == "optional":
            fc, ec, tc, ls, lw = "white", ORANGE, ORANGE, (0, (5, 3)), 2.0
        elif kind == "end":
            fc, ec, tc, ls, lw = GREEN, GREEN, "white", "-", 1.2
        else:
            fc, ec, tc, ls, lw = NAVY, NAVY, "white", "-", 1.2
        box = FancyBboxPatch(
            (x0, y),
            box_w,
            box_h,
            boxstyle="round,pad=0.02,rounding_size=0.12",
            facecolor=fc,
            edgecolor=ec,
            linewidth=lw,
            linestyle=ls,
        )
        ax.add_patch(box)
        title_y = y + 0.32 if step["sub"] else y + 0.48
        ax.text(
            x0 + box_w / 2,
            title_y,
            wrap_title(step["title"]),
            ha="center",
            va="center",
            fontsize=10.2,
            fontweight="bold",
            color=tc,
        )
        if step["sub"]:
            ax.text(
                x0 + box_w / 2,
                y + 0.68,
                step["sub"],
                ha="center",
                va="center",
                fontsize=8.2,
                color=tc,
                alpha=0.95,
            )
        centers.append((x0 + box_w / 2, y, y + box_h))
        if step.get("baca"):
            by = y + (box_h - baca_h) / 2
            baca = FancyBboxPatch(
                (baca_x, by),
                baca_w,
                baca_h,
                boxstyle="round,pad=0.02,rounding_size=0.10",
                facecolor=CHIMNEY_FILL,
                edgecolor=CHIMNEY_STROKE,
                linewidth=1.8,
            )
            ax.add_patch(baca)
            ax.text(
                baca_x + baca_w / 2,
                by + baca_h / 2,
                step["baca"],
                ha="center",
                va="center",
                fontsize=9.4,
                fontweight="bold",
                color=CHIMNEY_TEXT,
            )
            ax.add_patch(
                FancyArrowPatch(
                    (x0 + box_w + 0.04, y + box_h / 2),
                    (baca_x - 0.04, by + baca_h / 2),
                    arrowstyle="-|>",
                    mutation_scale=12,
                    linewidth=1.8,
                    color=CHIMNEY_STROKE,
                )
            )
        y += box_h + gap

    for i in range(len(centers) - 1):
        _, _y0, y_bottom = centers[i]
        _, y_top, _ = centers[i + 1]
        dashed = STEPS[i + 1]["kind"] == "optional"
        ax.add_patch(
            FancyArrowPatch(
                (centers[i][0], y_bottom + 0.02),
                (centers[i + 1][0], y_top - 0.02),
                arrowstyle="-|>",
                mutation_scale=13,
                linewidth=1.8,
                color=ORANGE if dashed else NAVY,
                linestyle=(0, (5, 3)) if dashed else "-",
            )
        )

    ax.text(
        0.55,
        13.95,
        "Lejant:  lacivert = proses   •   turuncu kesik = opsiyonel ODD   •   yeşil = sevk   •   krem kutu = emisyon bacası",
        ha="left",
        va="center",
        fontsize=8,
        color="#444444",
    )
    fig.tight_layout(pad=0.3)
    fig.savefig(path, dpi=160, facecolor="white")
    plt.close(fig)


def write_docx(png_path: Path, docx_path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(12)
    section.right_margin = Mm(12)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Boru Üretim Tesisi — Proses İş Akım Şeması")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    run.font.name = "Calibri"

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(6)
    nrun = note.add_run(
        "Winder, ODD tamir, kalibrasyon ve kaplin kanal açma adımlarına emisyon bacaları ok ile işlenmiştir."
    )
    nrun.font.size = Pt(9)
    nrun.font.name = "Calibri"

    usable_w = section.page_width - section.left_margin - section.right_margin
    usable_h = section.page_height - section.top_margin - section.bottom_margin - Mm(22)
    pic = document.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(png_path), height=usable_h)

    foot = document.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = foot.add_run(
        "Bacalar: Stiren Emiş Bacası 1 (Winder) · ODD Tamir Proses Bacası · "
        "Kalibrasyon Proses Bacası · Kaplin Kanal Açma Proses Bacası"
    )
    fr.font.size = Pt(8)
    fr.italic = True
    document.save(docx_path)


def main() -> None:
    drawio = OUT_DIR / "boru-uretim-is-akim-semasi.drawio"
    png = OUT_DIR / "boru-uretim-is-akim-semasi.png"
    docx = OUT_DIR / "boru-uretim-is-akim-semasi.docx"
    write_drawio(drawio)
    print(f"yazıldı: {drawio}")
    write_png(png)
    print(f"yazıldı: {png}")
    write_docx(png, docx)
    print(f"yazıldı: {docx}")


if __name__ == "__main__":
    main()
