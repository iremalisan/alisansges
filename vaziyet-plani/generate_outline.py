#!/usr/bin/env python3
"""B1 Blok vaziyet planı — yalnızca dış hatlar.

Fotoğraftaki vaziyet planından yaklaşık dış poligonlar üretir.
İç detay (çatı, ışıklık, eğim, kot) bilinçli olarak bırakılmaz.
Ölçüler şematiktir; resmi imar/ÇED paftası yerine geçmez.
"""

from __future__ import annotations

import math
from pathlib import Path
from xml.sax.saxutils import escape

DRAWIO_SCALE = 4.5  # 1 m → 4.5 px, draw.io’da rahat düzenleme

OUT_DIR = Path(__file__).resolve().parent
ROTATION_DEG = 30.0  # uzun cephe sağa-yukarı
MARGIN_M = 28.0

# Yerel metre: +X bina boyu (KD), +Y bina eni (KB). Dönüş sonrası kuzey yukarı kalır.
# Birleşik dikdörtgenler: batı üretim (geniş) + ara girinti + uzun depo.
BINA = [
    (0.0, 0.0),
    (175.0, 0.0),
    (175.0, 52.0),
    (82.0, 52.0),
    (82.0, 46.0),
    (70.0, 46.0),
    (70.0, 60.0),
    (22.0, 60.0),
    (22.0, 72.0),
    (10.0, 72.0),
    (10.0, 60.0),
    (0.0, 60.0),
    (0.0, 18.0),
    (-8.0, 18.0),
    (-8.0, 10.0),
    (0.0, 10.0),
    (0.0, 8.0),
    (-8.0, 8.0),
    (-8.0, 0.0),
]

# Tevsiiyat, binanın GD uzun cephesine bitişik; ortak duvar çizilmez.
TEVSIYAT = [
    (80.0, 0.0),
    (80.0, -56.0),
    (175.0, -56.0),
    (175.0, 0.0),
]

PARSEL = [
    (-48.0, 88.0),
    (22.0, 102.0),
    (96.0, 97.0),
    (158.0, 84.0),
    (198.0, 66.0),
    (210.0, 28.0),
    (206.0, -22.0),
    (198.0, -80.0),
    (138.0, -92.0),
    (68.0, -86.0),
    (8.0, -74.0),
    (-32.0, -42.0),
    (-54.0, 8.0),
    (-56.0, 52.0),
]

GIRISLER = [
    # (taban_orta, yön_içeri, etiket)
    ((128.0, 52.0), (0.0, -1.0), "DEPO GİRİŞİ"),
    ((175.0, 26.0), (-1.0, 0.0), "HAM MADDE\nBİNASI GİRİŞİ"),
]

ETIKETLER = [
    (88.0, 26.0, "B1 BLOK ANA ÜRETİM TESİSİ", 12, True),
    (127.0, -28.0, "TEVSİYAT ALANI", 11, True),
    (-20.0, 82.0, "STOK SAHASI", 10, True),
    (42.0, 90.0, "STOK SAHASI", 10, True),
    (16.0, 78.0, "ÜRÜN ÇIKIŞI", 7.5, True),
    (-20.0, 14.0, "ÜRÜN ÇIKIŞI", 7.5, True),
    (-20.0, 3.0, "ÜRÜN ÇIKIŞI", 7.5, True),
    (-22.0, -52.0, "208/5", 10, False),
]


def rect_poly(x0: float, y0: float, x1: float, y1: float) -> list[tuple[float, float]]:
    return [(x0, y0), (x1, y0), (x1, y1), (x0, y1)]


# draw.io içinde sürüklenebilir örnek hacimler (silinebilir / kopyalanabilir)
IC_HACIMLER = [
    (rect_poly(6.0, 8.0, 66.0, 54.0), "ÜRETİM", "#dae8fc", "#6c8ebf"),
    (rect_poly(84.0, 8.0, 168.0, 46.0), "DEPO", "#d5e8d4", "#82b366"),
    (rect_poly(90.0, -50.0, 168.0, -8.0), "TEVSİYAT / GENİŞLEME", "#fff2cc", "#d6b656"),
]


def rot(x: float, y: float) -> tuple[float, float]:
    a = math.radians(ROTATION_DEG)
    return (
        x * math.cos(a) - y * math.sin(a),
        x * math.sin(a) + y * math.cos(a),
    )


def rot_pts(pts: list[tuple[float, float]]) -> list[tuple[float, float]]:
    return [rot(x, y) for x, y in pts]


def bounds(groups: list[list[tuple[float, float]]]) -> tuple[float, float, float, float]:
    xs = [p[0] for g in groups for p in g]
    ys = [p[1] for g in groups for p in g]
    return min(xs), min(ys), max(xs), max(ys)


def svg_path(pts: list[tuple[float, float]], closed: bool) -> str:
    d = [f"M {pts[0][0]:.3f},{pts[0][1]:.3f}"]
    d += [f"L {x:.3f},{y:.3f}" for x, y in pts[1:]]
    if closed:
        d.append("Z")
    return " ".join(d)


def triangle(
    origin: tuple[float, float], direction: tuple[float, float], size: float = 4.2
) -> list[tuple[float, float]]:
    dx, dy = direction
    length = math.hypot(dx, dy) or 1.0
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    tip = (origin[0] + ux * size, origin[1] + uy * size)
    left = (origin[0] + px * size * 0.45, origin[1] + py * size * 0.45)
    right = (origin[0] - px * size * 0.45, origin[1] - py * size * 0.45)
    return [tip, left, right]


def write_svg(
    bina: list[tuple[float, float]],
    tevsiyat: list[tuple[float, float]],
    parsel: list[tuple[float, float]],
    giris_tris: list[list[tuple[float, float]]],
    labels: list[tuple[float, float, str, float, bool]],
    minx: float,
    miny: float,
    maxx: float,
    maxy: float,
    path: Path,
) -> None:
    width = maxx - minx
    height = maxy - miny

    def sy(y: float) -> float:
        return maxy - y  # SVG'de y aşağı artar; kuzey yukarı kalsın

    def sp(pts: list[tuple[float, float]]) -> list[tuple[float, float]]:
        return [(x, sy(y)) for x, y in pts]

    frame = [
        (minx + 6, miny + 6),
        (maxx - 6, miny + 6),
        (maxx - 6, maxy - 6),
        (minx + 6, maxy - 6),
    ]
    north_base = (maxx - 38, maxy - 36)
    scale_origin = (minx + 22, miny + 18)

    svg_labels = []
    for x, y, text, size, follow in labels:
        angle = -ROTATION_DEG if follow else 0.0
        for i, line in enumerate(text.split("\n")):
            svg_labels.append((x, sy(y) + i * size * 1.15, line, size, angle))

    title_x, title_y = maxx - 18, sy(miny + 22)
    note_x, note_y = minx + 18, sy(maxy - 14)

    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="{minx:.2f} {sy(maxy):.2f} {width:.2f} {height:.2f}" '
        f'width="1600" height="{1600 * height / width:.0f}">',
        "  <title>B1 Blok Ana Üretim Tesisi — vaziyet planı dış hat</title>",
        "  <desc>Yalnızca parsel, bina ve tevsiiyat dış çizgileri. İç dolgu için boş bırakılmıştır.</desc>",
        '  <rect x="{:.2f}" y="{:.2f}" width="{:.2f}" height="{:.2f}" fill="#ffffff"/>'.format(
            minx, sy(maxy), width, height
        ),
        '  <g id="cerceve" fill="none" stroke="#111111" stroke-width="0.6">',
        f'    <path d="{svg_path(sp(frame), True)}"/>',
        "  </g>",
        '  <g id="parsel" fill="none" stroke="#111111" stroke-width="1.15" '
        'stroke-dasharray="7 3.5" stroke-linejoin="miter">',
        f'    <path d="{svg_path(sp(parsel), True)}"/>',
        "  </g>",
        '  <g id="tevsiyat" fill="none" stroke="#111111" stroke-width="0.95" '
        'stroke-dasharray="5 3" stroke-linejoin="miter">',
        f'    <path d="{svg_path(sp(tevsiyat), False)}"/>',
        "  </g>",
        '  <g id="bina" fill="none" stroke="#111111" stroke-width="1.35" stroke-linejoin="miter">',
        f'    <path d="{svg_path(sp(bina), True)}"/>',
        "  </g>",
        '  <g id="girisler" fill="#111111" stroke="none">',
    ]
    for tri in giris_tris:
        parts.append(f'    <path d="{svg_path(sp(tri), True)}"/>')
    parts.append("  </g>")
    parts.append('  <g id="yazi" fill="#111111" font-family="Arial, Helvetica, sans-serif">')
    for x, y, text, size, angle in svg_labels:
        parts.append(
            f'    <text x="{x:.2f}" y="{y:.2f}" font-size="{size}" text-anchor="middle" '
            f'dominant-baseline="middle" transform="rotate({angle:.1f} {x:.2f} {y:.2f})">{text}</text>'
        )
    parts.append(
        f'    <text x="{title_x:.2f}" y="{title_y:.2f}" font-size="7.5" text-anchor="end">'
        "VAZİYET PLANI — DIŞ HAT  ·  şematik  ·  içini doldurmak için boş</text>"
    )
    parts.append(
        f'    <text x="{note_x:.2f}" y="{note_y:.2f}" font-size="6.2" text-anchor="start">'
        "Fotoğraftan yaklaşık dış hat. Resmi pafta / kot / koordinat içermez.</text>"
    )
    parts.append("  </g>")

    # Kuzey oku (kuzey yukarı)
    nx, ny = north_base[0], sy(north_base[1])
    parts.append('  <g id="kuzey" fill="#111111" stroke="#111111" stroke-width="0.7">')
    parts.append(
        f'    <line x1="{nx:.2f}" y1="{ny + 10:.2f}" x2="{nx:.2f}" y2="{ny - 14:.2f}"/>'
    )
    parts.append(
        f'    <path d="M {nx:.2f},{ny - 16:.2f} L {nx - 3.2:.2f},{ny - 8:.2f} L {nx + 3.2:.2f},{ny - 8:.2f} Z"/>'
    )
    parts.append(
        f'    <text x="{nx:.2f}" y="{ny - 20:.2f}" font-size="8" text-anchor="middle" '
        'font-family="Arial, Helvetica, sans-serif">K</text>'
    )
    parts.append("  </g>")

    # Ölçek çubuğu 0–50 m
    sx, syb = scale_origin[0], sy(scale_origin[1])
    parts.append('  <g id="olcek" fill="none" stroke="#111111" stroke-width="0.8">')
    parts.append(f'    <line x1="{sx:.2f}" y1="{syb:.2f}" x2="{sx + 50:.2f}" y2="{syb:.2f}"/>')
    for t, label in ((0, "0"), (25, "25"), (50, "50 m")):
        parts.append(
            f'    <line x1="{sx + t:.2f}" y1="{syb - 2:.2f}" x2="{sx + t:.2f}" y2="{syb + 2:.2f}"/>'
        )
        parts.append(
            f'    <text x="{sx + t:.2f}" y="{syb + 9:.2f}" font-size="5.5" text-anchor="middle" '
            f'fill="#111111" stroke="none" font-family="Arial, Helvetica, sans-serif">{label}</text>'
        )
    parts.append("  </g>")
    parts.append("</svg>")
    path.write_text("\n".join(parts) + "\n", encoding="utf-8")


def dxf_escape(text: str) -> str:
    return text.replace("\n", " ")


def write_dxf(
    bina: list[tuple[float, float]],
    tevsiyat: list[tuple[float, float]],
    parsel: list[tuple[float, float]],
    giris_tris: list[list[tuple[float, float]]],
    labels: list[tuple[float, float, str, float, bool]],
    minx: float,
    miny: float,
    maxx: float,
    maxy: float,
    path: Path,
) -> None:
    lines: list[str] = []

    def pair(code: int, value: object) -> None:
        lines.append(str(code))
        lines.append(str(value))

    def polyline(pts: list[tuple[float, float]], layer: str, closed: bool, ltype: str = "CONTINUOUS") -> None:
        pair(0, "POLYLINE")
        pair(8, layer)
        pair(6, ltype)
        pair(66, 1)
        pair(70, 1 if closed else 0)
        for x, y in pts:
            pair(0, "VERTEX")
            pair(8, layer)
            pair(10, f"{x:.4f}")
            pair(20, f"{y:.4f}")
            pair(30, "0.0")
        pair(0, "SEQEND")
        pair(8, layer)

    def solid(pts: list[tuple[float, float]], layer: str) -> None:
        pair(0, "SOLID")
        pair(8, layer)
        pair(10, f"{pts[0][0]:.4f}")
        pair(20, f"{pts[0][1]:.4f}")
        pair(30, "0.0")
        pair(11, f"{pts[1][0]:.4f}")
        pair(21, f"{pts[1][1]:.4f}")
        pair(31, "0.0")
        pair(12, f"{pts[2][0]:.4f}")
        pair(22, f"{pts[2][1]:.4f}")
        pair(32, "0.0")
        pair(13, f"{pts[2][0]:.4f}")
        pair(23, f"{pts[2][1]:.4f}")
        pair(33, "0.0")

    def text(x: float, y: float, content: str, height: float, layer: str, rotation: float = 0.0) -> None:
        pair(0, "TEXT")
        pair(8, layer)
        pair(10, f"{x:.4f}")
        pair(20, f"{y:.4f}")
        pair(30, "0.0")
        pair(40, f"{height:.3f}")
        pair(1, dxf_escape(content))
        pair(50, f"{rotation:.3f}")
        pair(72, 1)
        pair(11, f"{x:.4f}")
        pair(21, f"{y:.4f}")
        pair(31, "0.0")

    pair(0, "SECTION")
    pair(2, "HEADER")
    pair(9, "$ACADVER")
    pair(1, "AC1009")
    pair(9, "$INSUNITS")
    pair(70, 6)  # metre
    pair(0, "ENDSEC")

    pair(0, "SECTION")
    pair(2, "TABLES")
    pair(0, "TABLE")
    pair(2, "LTYPE")
    pair(70, 2)
    pair(0, "LTYPE")
    pair(2, "CONTINUOUS")
    pair(70, 0)
    pair(3, "Solid line")
    pair(72, 65)
    pair(73, 0)
    pair(40, "0.0")
    pair(0, "LTYPE")
    pair(2, "DASHED")
    pair(70, 0)
    pair(3, "Dashed")
    pair(72, 65)
    pair(73, 2)
    pair(40, "10.0")
    pair(49, "7.0")
    pair(49, "-3.0")
    pair(0, "ENDTAB")
    pair(0, "TABLE")
    pair(2, "LAYER")
    pair(70, 5)
    for name, color, ltype in (
        ("PARSEL", 7, "DASHED"),
        ("BINA", 7, "CONTINUOUS"),
        ("TEVSIYAT", 7, "DASHED"),
        ("GIRIS", 7, "CONTINUOUS"),
        ("YAZI", 7, "CONTINUOUS"),
    ):
        pair(0, "LAYER")
        pair(2, name)
        pair(70, 0)
        pair(62, color)
        pair(6, ltype)
    pair(0, "ENDTAB")
    pair(0, "ENDSEC")

    pair(0, "SECTION")
    pair(2, "ENTITIES")
    polyline(parsel, "PARSEL", True, "DASHED")
    polyline(bina, "BINA", True, "CONTINUOUS")
    polyline(tevsiyat, "TEVSIYAT", False, "DASHED")
    for tri in giris_tris:
        solid(tri, "GIRIS")
    for x, y, content, size, follow in labels:
        angle = ROTATION_DEG if follow else 0.0
        for i, line in enumerate(content.split("\n")):
            text(x, y - i * size * 1.1, line, size * 0.85, "YAZI", angle)
    text(maxx - 8, miny + 10, "VAZIYET PLANI - DIS HAT", 3.2, "YAZI", 0.0)
    text(minx + 8, maxy - 8, "K", 4.0, "YAZI", 0.0)
    # Kuzey oku
    pair(0, "LINE")
    pair(8, "YAZI")
    pair(10, f"{maxx - 38:.4f}")
    pair(20, f"{maxy - 22:.4f}")
    pair(11, f"{maxx - 38:.4f}")
    pair(21, f"{maxy - 48:.4f}")
    north_tri = triangle((maxx - 38, maxy - 48), (0.0, 1.0), 6.0)
    solid(north_tri, "YAZI")
    pair(0, "ENDSEC")
    pair(0, "EOF")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_png(
    bina: list[tuple[float, float]],
    tevsiyat: list[tuple[float, float]],
    parsel: list[tuple[float, float]],
    giris_tris: list[list[tuple[float, float]]],
    labels: list[tuple[float, float, str, float, bool]],
    minx: float,
    miny: float,
    maxx: float,
    maxy: float,
    path: Path,
) -> None:
    import matplotlib.pyplot as plt
    from matplotlib.patches import Polygon

    fig_w = 14
    fig_h = fig_w * (maxy - miny) / (maxx - minx)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=160)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    def add_poly(pts: list[tuple[float, float]], closed: bool, **kwargs) -> None:
        xs = [p[0] for p in pts] + ([pts[0][0]] if closed else [])
        ys = [p[1] for p in pts] + ([pts[0][1]] if closed else [])
        ax.plot(xs, ys, **kwargs)

    add_poly(parsel, True, color="#111111", lw=1.6, ls=(0, (7, 3.5)), solid_capstyle="butt")
    add_poly(tevsiyat, False, color="#111111", lw=1.4, ls=(0, (5, 3)), solid_capstyle="butt")
    add_poly(bina, True, color="#111111", lw=2.0, ls="-")
    for tri in giris_tris:
        ax.add_patch(Polygon(tri, closed=True, facecolor="#111111", edgecolor="none"))

    for x, y, content, size, follow in labels:
        ax.text(
            x,
            y,
            content,
            fontsize=size * 0.72,
            ha="center",
            va="center",
            color="#111111",
            rotation=ROTATION_DEG if follow else 0.0,
            linespacing=1.15,
        )

    ax.annotate(
        "",
        xy=(maxx - 38, maxy - 18),
        xytext=(maxx - 38, maxy - 42),
        arrowprops=dict(arrowstyle="-|>", color="#111111", lw=1.2),
    )
    ax.text(maxx - 38, maxy - 14, "K", ha="center", va="bottom", fontsize=9, color="#111111")
    ax.plot([minx + 22, minx + 72], [miny + 16, miny + 16], color="#111111", lw=1.2)
    ax.text(minx + 22, miny + 12, "0", ha="center", va="top", fontsize=7)
    ax.text(minx + 72, miny + 12, "50 m", ha="center", va="top", fontsize=7)
    ax.text(
        maxx - 10,
        miny + 10,
        "VAZİYET PLANI — DIŞ HAT  ·  içini doldurmak için boş",
        ha="right",
        va="bottom",
        fontsize=7.5,
        color="#111111",
    )
    ax.set_aspect("equal")
    ax.set_xlim(minx, maxx)
    ax.set_ylim(miny, maxy)
    ax.axis("off")
    fig.tight_layout(pad=0.2)
    fig.savefig(path, dpi=180, facecolor="white")
    plt.close(fig)


def write_bmp(png_path: Path, bmp_path: Path) -> None:
    from PIL import Image

    image = Image.open(png_path).convert("RGB")
    image.save(bmp_path, format="BMP")


def write_docx(png_path: Path, docx_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    from PIL import Image

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(420)
    section.page_height = Mm(297)
    section.left_margin = Mm(12)
    section.right_margin = Mm(12)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(12)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("VAZİYET PLANI — DIŞ HAT")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    run.font.name = "Calibri"

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    sub = subtitle.add_run(
        "B1 Blok Ana Üretim Tesisi  ·  içini Word’de şekil ekleyerek doldurabilirsiniz"
    )
    sub.font.size = Pt(11)
    sub.font.name = "Calibri"

    usable_width = section.page_width - section.left_margin - section.right_margin
    usable_height = (
        section.page_height
        - section.top_margin
        - section.bottom_margin
        - Mm(32)
    )
    with Image.open(png_path) as preview:
        aspect = preview.height / preview.width
    width = usable_width
    height = int(width * aspect)
    if height > usable_height:
        width = int(usable_height / aspect)

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(6)
    picture.add_run().add_picture(str(png_path), width=width)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note_run = note.add_run(
        "Not: Fotoğraftan yaklaşık dış hat. Resmi imar / ÇED / ruhsat paftası yerine geçmez. "
        "Paint ile boyamak için aynı klasördeki .bmp dosyasını açın."
    )
    note_run.font.size = Pt(9)
    note_run.italic = True
    note_run.font.name = "Calibri"

    document.save(docx_path)


def _canvas_pts(
    pts: list[tuple[float, float]], minx: float, maxy: float, scale: float
) -> list[tuple[float, float]]:
    return [((x - minx) * scale, (maxy - y) * scale) for x, y in pts]


def _poly_coords(pts: list[tuple[float, float]]) -> tuple[float, float, float, float, str]:
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    width = max(x1 - x0, 1.0)
    height = max(y1 - y0, 1.0)
    coords = ",".join(f"[{(x - x0) / width:.5f},{(y - y0) / height:.5f}]" for x, y in pts)
    return x0, y0, width, height, coords


def _drawio_value(text: str) -> str:
    return escape(text).replace("\n", "&lt;br&gt;")


def write_drawio(
    bina: list[tuple[float, float]],
    tevsiyat: list[tuple[float, float]],
    parsel: list[tuple[float, float]],
    giris_tris: list[list[tuple[float, float]]],
    labels: list[tuple[float, float, str, float, bool]],
    ic_hacimler: list[tuple[list[tuple[float, float]], str, str, str]],
    minx: float,
    miny: float,
    maxx: float,
    maxy: float,
    path: Path,
) -> None:
    scale = DRAWIO_SCALE
    page_w = (maxx - minx) * scale
    page_h = (maxy - miny) * scale
    bina_c = _canvas_pts(bina, minx, maxy, scale)
    tev_c = _canvas_pts(tevsiyat, minx, maxy, scale)
    parsel_c = _canvas_pts(parsel, minx, maxy, scale)
    tris_c = [_canvas_pts(tri, minx, maxy, scale) for tri in giris_tris]
    labels_c = [
        (((x - minx) * scale, (maxy - y) * scale, text, size, follow))
        for x, y, text, size, follow in labels
    ]
    rooms_c = [
        (_canvas_pts(poly, minx, maxy, scale), name, fill, stroke)
        for poly, name, fill, stroke in ic_hacimler
    ]

    cells: list[str] = [
        '    <mxCell id="0"/>',
        '    <mxCell id="1" parent="0"/>',
        '    <mxCell id="parsel" value="Parsel" parent="0"/>',
        '    <mxCell id="tevsiyat" value="Tevsiiyat" parent="0"/>',
        '    <mxCell id="bina" value="Bina" parent="0"/>',
        '    <mxCell id="ic" value="Ic hacimler (surukle / kopyala / boya)" parent="0"/>',
        '    <mxCell id="giris" value="Girisler" parent="0"/>',
        '    <mxCell id="yazi" value="Yazilar" parent="0"/>',
        '    <mxCell id="sembol" value="Semboller" parent="0"/>',
    ]

    def add_polygon(
        cell_id: str,
        parent: str,
        pts: list[tuple[float, float]],
        extra_style: str,
        value: str = "",
        closed: bool = True,
    ) -> None:
        x0, y0, width, height, coords = _poly_coords(pts)
        style = (
            "html=1;whiteSpace=wrap;shape=mxgraph.basic.polygon;"
            f"polyCoords=[{coords}];polyline={0 if closed else 1};"
            "movable=1;resizable=1;rotatable=1;deletable=1;editable=1;"
            f"{extra_style}"
        )
        cells.append(
            f'    <mxCell id="{cell_id}" value="{_drawio_value(value)}" style="{style}" '
            f'vertex="1" parent="{parent}">\n'
            f'      <mxGeometry x="{x0:.2f}" y="{y0:.2f}" width="{width:.2f}" '
            f'height="{height:.2f}" as="geometry"/>\n'
            "    </mxCell>"
        )

    def add_text(
        cell_id: str,
        parent: str,
        x: float,
        y: float,
        text: str,
        size: float,
        follow: bool,
    ) -> None:
        font = max(11.0, size * 1.55)
        lines = text.split("\n")
        longest = max(len(line) for line in lines)
        width = max(80.0, longest * font * 0.72)
        height = font * 1.7 * len(lines) + 8
        rotation = f"rotation={-ROTATION_DEG};" if follow else ""
        style = (
            "text;html=1;align=center;verticalAlign=middle;whiteSpace=wrap;"
            f"fontFamily=Arial;fontSize={font:.1f};fontColor=#111111;"
            "movable=1;resizable=1;rotatable=1;deletable=1;editable=1;"
            f"{rotation}labelBackgroundColor=none;"
        )
        cells.append(
            f'    <mxCell id="{cell_id}" value="{_drawio_value(text)}" style="{style}" '
            f'vertex="1" parent="{parent}">\n'
            f'      <mxGeometry x="{x - width / 2:.2f}" y="{y - height / 2:.2f}" '
            f'width="{width:.2f}" height="{height:.2f}" as="geometry"/>\n'
            "    </mxCell>"
        )

    add_polygon(
        "parsel-hat",
        "parsel",
        parsel_c,
        "fillColor=none;strokeColor=#111111;strokeWidth=2;dashed=1;dashPattern=12 6;",
    )
    add_polygon(
        "tevsiyat-hat",
        "tevsiyat",
        tev_c,
        "fillColor=#f5f5f5;strokeColor=#111111;strokeWidth=2;dashed=1;dashPattern=8 4;",
    )
    add_polygon(
        "bina-hat",
        "bina",
        bina_c,
        "fillColor=#ffffff;strokeColor=#111111;strokeWidth=2.6;",
    )
    for i, (poly, name, fill, stroke) in enumerate(rooms_c, start=1):
        add_polygon(
            f"oda-{i}",
            "ic",
            poly,
            f"fillColor={fill};strokeColor={stroke};strokeWidth=1.5;opacity=75;fontSize=16;"
            "fontFamily=Arial;fontStyle=1;align=center;verticalAlign=middle;",
            value=name,
        )
    for i, tri in enumerate(tris_c, start=1):
        add_polygon(
            f"giris-{i}",
            "giris",
            tri,
            "fillColor=#111111;strokeColor=#111111;strokeWidth=1;",
        )
    for i, (x, y, text, size, follow) in enumerate(labels_c, start=1):
        add_text(f"yazi-{i}", "yazi", x, y, text, size, follow)

    # Kuzey oku
    nx, ny = page_w - 90, 70
    cells.append(
        f'    <mxCell id="kuzey-ok" value="" style="endArrow=block;endFill=1;html=1;'
        f'strokeColor=#111111;strokeWidth=2;movable=1;resizable=1;rotatable=1;" '
        f'edge="1" parent="sembol">\n'
        f'      <mxGeometry relative="1" as="geometry">\n'
        f'        <mxPoint x="{nx:.2f}" y="{ny + 36:.2f}" as="sourcePoint"/>\n'
        f'        <mxPoint x="{nx:.2f}" y="{ny - 20:.2f}" as="targetPoint"/>\n'
        f"      </mxGeometry>\n"
        "    </mxCell>"
    )
    add_text("kuzey-k", "sembol", nx, ny - 36, "K", 12, False)

    # 50 m ölçek
    sx, syb = 80.0, page_h - 55
    cells.append(
        f'    <mxCell id="olcek" value="" style="endArrow=none;startArrow=none;html=1;'
        f'strokeColor=#111111;strokeWidth=2;movable=1;" edge="1" parent="sembol">\n'
        f'      <mxGeometry relative="1" as="geometry">\n'
        f'        <mxPoint x="{sx:.2f}" y="{syb:.2f}" as="sourcePoint"/>\n'
        f'        <mxPoint x="{sx + 50 * scale:.2f}" y="{syb:.2f}" as="targetPoint"/>\n'
        f"      </mxGeometry>\n"
        "    </mxCell>"
    )
    add_text("olcek-0", "sembol", sx, syb + 18, "0", 8, False)
    add_text("olcek-50", "sembol", sx + 50 * scale, syb + 18, "50 m", 8, False)

    add_text(
        "baslik",
        "yazi",
        page_w - 220,
        page_h - 28,
        "VAZİYET PLANI — DRAW.IO  ·  şekilleri sürükle, boya, kopyala",
        8,
        False,
    )
    cells.append(
        '    <mxCell id="not" value="'
        + _drawio_value(
            "Nasıl oynarım?\n"
            "• Şekli tutup sürükle\n"
            "• Çift tıkla yazıyı değiştir\n"
            "• Fill / Line ile boya\n"
            "• Ctrl+D ile oda kopyala\n"
            "• Soldan dikdörtgen ekle"
        )
        + '" style="shape=note;whiteSpace=wrap;html=1;align=left;verticalAlign=top;'
        "fillColor=#fff2cc;strokeColor=#d6b656;fontSize=12;fontFamily=Arial;"
        'spacing=8;movable=1;resizable=1;deletable=1;editable=1;" vertex="1" parent="sembol">\n'
        f'      <mxGeometry x="16" y="16" width="230" height="150" as="geometry"/>\n'
        "    </mxCell>"
    )

    xml = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        '<mxfile host="app.diagrams.net" agent="GES-Metraj-Pro" version="22.1.0">',
        '  <diagram id="b1-vaziyet" name="B1 Blok Vaziyet Dış Hat">',
        f'    <mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" '
        f'connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="{page_w:.0f}" '
        f'pageHeight="{page_h:.0f}" math="0" shadow="0">',
        "      <root>",
        *cells,
        "      </root>",
        "    </mxGraphModel>",
        "  </diagram>",
        "</mxfile>",
        "",
    ]
    path.write_text("\n".join(xml), encoding="utf-8")


def main() -> None:
    bina = rot_pts(BINA)
    tevsiyat = rot_pts(TEVSIYAT)
    parsel = rot_pts(PARSEL)
    giris_tris = []
    for origin, direction, _label in GIRISLER:
        o = rot(*origin)
        d = rot(*direction)
        # yön vektörünü de döndür (orijin farkı)
        d0 = rot(0.0, 0.0)
        d = (d[0] - d0[0], d[1] - d0[1])
        giris_tris.append(triangle(o, d, 4.4))

    labels = []
    for x, y, text, size, follow in ETIKETLER:
        xr, yr = rot(x, y)
        labels.append((xr, yr, text, size, follow))
    for origin, _direction, text in GIRISLER:
        ox, oy = origin
        dx, dy = _direction
        # etiketi girişin dışına koy
        lx, ly = rot(ox - dx * 12.0, oy - dy * 12.0)
        labels.append((lx, ly, text, 7.5, True))

    minx, miny, maxx, maxy = bounds([bina, tevsiyat, parsel])
    minx -= MARGIN_M
    miny -= MARGIN_M
    maxx += MARGIN_M
    maxy += MARGIN_M

    svg_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.svg"
    dxf_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.dxf"
    png_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.png"
    bmp_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.bmp"
    docx_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.docx"
    drawio_path_out = OUT_DIR / "b1-blok-vaziyet-dis-hat.drawio"
    ic_hacimler = [
        (rot_pts(poly), name, fill, stroke) for poly, name, fill, stroke in IC_HACIMLER
    ]
    write_svg(bina, tevsiyat, parsel, giris_tris, labels, minx, miny, maxx, maxy, svg_path_out)
    write_dxf(bina, tevsiyat, parsel, giris_tris, labels, minx, miny, maxx, maxy, dxf_path_out)
    write_drawio(
        bina,
        tevsiyat,
        parsel,
        giris_tris,
        labels,
        ic_hacimler,
        minx,
        miny,
        maxx,
        maxy,
        drawio_path_out,
    )
    print(f"yazıldı: {svg_path_out}")
    print(f"yazıldı: {dxf_path_out}")
    print(f"yazıldı: {drawio_path_out}")
    try:
        write_png(bina, tevsiyat, parsel, giris_tris, labels, minx, miny, maxx, maxy, png_path_out)
        print(f"yazıldı: {png_path_out}")
        write_bmp(png_path_out, bmp_path_out)
        print(f"yazıldı: {bmp_path_out}")
        write_docx(png_path_out, docx_path_out)
        print(f"yazıldı: {docx_path_out}")
    except ImportError as exc:
        print(f"PNG/BMP/Word atlandı: {exc}")


if __name__ == "__main__":
    main()
