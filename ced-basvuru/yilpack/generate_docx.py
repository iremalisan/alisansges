#!/usr/bin/env python3
"""Yılpack ÇED proje özeti ve iş akım şeması — düzenlenebilir Word çıktısı."""

from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

PDF_PATH = Path("/workspace/ced-basvuru/yilpack/YILPACK_Proje_Ozeti_ve_Is_Akim_Semasi.pdf")
DOCX_PATH = Path("/workspace/ced-basvuru/yilpack/YILPACK_Proje_Ozeti_ve_Is_Akim_Semasi.docx")
FLOW_PNG = Path("/tmp/yilpack-is-akim-semasi.png")

TIMES = "Times New Roman"


def set_run_font(run, name=TIMES, size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6,
                         first_line=0, line=15):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(line)
    pf.first_line_indent = Cm(first_line)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY


def add_title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=12, line=16)
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)
    return p


def add_heading_text(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, line=16)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True)
    return p


def add_body(doc, text, *, first_line=0.5, space_after=6):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=first_line, space_after=space_after, line=16)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_subheading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=6, space_after=2, line=16)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "7F7F7F")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_text(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(12)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def export_flowchart_png():
    pdf = pymupdf.open(PDF_PATH)
    page = pdf[1]
    clip = pymupdf.Rect(36, 108, 560, 655)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), clip=clip, alpha=False)
    pix.save(str(FLOW_PNG))
    pdf.close()


def build():
    export_flowchart_png()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)

    style = doc.styles["Normal"]
    style.font.name = TIMES
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), TIMES)

    add_title(doc, "PROJE ÖZETİ")

    add_heading_text(doc, "1. TESİSİN YERİ")
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=0.5, space_after=6, line=16)
    run = p.add_run("TESİS ADRESİNİ YAZINIZ")
    set_run_font(run, size=11, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run = p.add_run(
        " adresinde, “Polietilen Poşet, Rafya İplik, Dokuma, Laminasyon "
        "(Kömür Torbası) ve Gübre Torbası Üretimi” faaliyetleri yapılması planlanmaktadır."
    )
    set_run_font(run, size=11)

    add_heading_text(doc, "2. PROJENİN TÜRÜ")
    add_body(
        doc,
        "Polietilen poşet, rafya iplik, dokuma, laminasyon (kömür torbaları) ve gübre "
        "torbası üretim faaliyetleri yapılması planlanmaktadır. Aşağıda proses özeti belirtilmiştir.",
    )
    add_subheading(doc, "Polietilen Poşet, Rafya İplik, Dokuma ve Torba Üretimi;")
    add_body(
        doc,
        "Firmanın işyerinde polietilen poşet, rafya iplik, dokuma kumaş, kömür torbası "
        "ve gübre torbası üretimi yapılmaktadır. Tesise gelen hammaddeler ilk olarak "
        "kantardan geçmekte olup tartım işlemi yapılmaktadır. Tartım işleminden sonra "
        "hammaddeler extruder hattında rafya iplik üretimine alınmaktadır. Üretilen iplikler "
        "dokuma ünitesinde kumaş haline getirilmektedir.",
    )
    add_body(
        doc,
        "Dokuma sonrası ürünler polietilen poşet üretimi, laminasyon (kömür torbaları) "
        "ve baskısız kumaş olarak konfeksiyon hatlarına yönlendirilmektedir. Laminasyon "
        "sonrası ürünler baskılı veya baskısız olarak konfeksiyon işlemine alınmaktadır. "
        "Konfeksiyon sonrası paketleme / ambalaj işlemi yapılmakta; gübre torbası için "
        "iç geçirme işlemi uygulanan ürünler ile birlikte sevkiyat gerçekleştirilmektedir.",
    )

    add_heading_text(doc, "3. TESİSTE KAYNAKLANACAK SIVI ATIKLAR")
    add_body(
        doc,
        "Tesiste çalışan personelden dolayı oluşacak atıksular, Belediyeye ait kanalizasyon "
        "sistemine verilmektedir. Tesiste endüstriyel nitelikli herhangi bir atıksu oluşmayacaktır.",
    )

    add_heading_text(doc, "4. TESİSTE KAYNAKLANACAK EMİSYONLAR")
    add_body(
        doc,
        "Tesiste ısınma ve sanayi amaçlı olarak elektrik enerjisi kullanılacaktır. Tesiste herhangi bir "
        "emisyon kaynağı olduğu taktirde SKHKKY hükümleri çerçevesinde gerekli ölçümler yaptırılıp söz "
        "konusu yönetmelikte belirtilen sınır değerleri sağlanacak, Çevre ve Şehircilik İl Müdürlüğü’ne bilgi "
        "verilecektir.",
    )

    add_heading_text(doc, "5. TESİSTE KAYNAKLANACAK ATIKLAR")
    add_body(doc, "Tesiste oluşabilecek atık kodları", first_line=0, space_after=4)

    rows = [
        ("ATIK KODU", "ATIK KODU TANIMI"),
        ("15 01 10*", "Tehlikeli maddelerin kalıntılarını içeren ya da tehlikeli maddelerle kontamine olmuş ambalajlar"),
        ("15 02 02*", "Tehlikeli maddelerle kirlenmiş emiciler, filtre malzemeleri (başka şekilde tanımlanmamış ise yağ filtreleri), temizleme bezleri, koruyucu giysiler"),
        ("20 01 21*", "Fluoresan lambalar ve diğer cıva içeren atıklar"),
        ("20 01 08", "Biyolojik olarak bozunabilir mutfak ve kantin atıkları"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (kod, tanim) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(3.2)
        c1.width = Cm(13.8)
        if i == 0:
            set_cell_text(c0, kod, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            set_cell_text(c1, tanim, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            shade_cell(c0, "D9E2F3")
            shade_cell(c1, "D9E2F3")
        else:
            set_cell_text(c0, kod, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
            set_cell_text(c1, tanim, size=9)
            if i % 2 == 0:
                shade_cell(c0, "F2F2F2")
                shade_cell(c1, "F2F2F2")
        set_cell_border(c0)
        set_cell_border(c1)

    p = doc.add_paragraph()
    set_paragraph_format(p, space_before=8, space_after=6, first_line=0.5, line=16)
    run = p.add_run(
        "Tehlikeli olmayan atıklar, Çevre, Şehircilik ve İklim Değişikliği Bakanlığı’ndan yetki almış "
        "geri dönüşüm/kazanım firmalarına verilecektir. Oluşabilecek tehlikeli atıklar ise, Çevre, Şehircilik ve "
        "İklim Değişikliği Bakanlığı’ndan yetki almış bertaraf/geri kazanım tesislerine gönderilecek ve "
        "bertarafı sağlanacaktır."
    )
    set_run_font(run, size=11)

    add_body(
        doc,
        "Tesiste oluşacak evsel atıkların bertarafı, Belediye tarafından sağlanacaktır.",
    )

    add_heading_text(doc, "6. TESİSTE KAYNAKLANACAK GÜRÜLTÜ")
    add_body(doc, "Tesiste herhangi bir gürültü kaynağı mevcut değildir.")

    doc.add_page_break()

    for text, size in (
        ("YILPACK", 12),
        ("POLİETİLEN POŞET, RAFYA İPLİK, DOKUMA VE TORBA ÜRETİM TESİSİ", 12),
        ("İŞ AKIM ŞEMASI", 12),
    ):
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, line=16)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=True)

    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=0, line=12)
    p.add_run().add_picture(str(FLOW_PNG), width=Mm(168))

    doc.save(str(DOCX_PATH))
    print("Wrote", DOCX_PATH)


if __name__ == "__main__":
    build()
